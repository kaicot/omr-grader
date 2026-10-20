"""A 40- or 60-question form gets books with exactly that many Q columns."""

from __future__ import annotations

from dataclasses import replace
from hashlib import sha256
from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import Workbook, load_workbook

from omr_grader.application.dto import AnswerKeyValidation, ScoreInput
from omr_grader.application.grading_presenter import GradingPresenter
from omr_grader.application.grading_use_case import GradingUseCase
from omr_grader.domain.enums import (
    AnswerKeySnapshotKind,
    AnswerStatus,
    KeyQuestionStatus,
    SourceKind,
    StudentIdStatus,
)
from omr_grader.domain.errors import Err, Ok
from omr_grader.domain.grading import score_effective
from omr_grader.domain.models import (
    AnswerKeyEntry,
    AnswerKeySnapshot,
    AnswerValue,
    EffectiveResponse,
    ImportedResponseRef,
)
from omr_grader.workbooks.response_book import write_response_projection
from omr_grader.workbooks.response_import import parse_response_book
from omr_grader.workbooks.schemas import (
    RESPONSE_FIXED_HEADERS,
    RESPONSE_SHEET_NAME,
    question_count,
    question_headers,
    response_headers,
)
from omr_grader.workbooks.score_book import (
    book_question_count,
    write_final_book,
    write_score_book,
)

UNASKED = AnswerValue((), AnswerStatus.UNASKED)
SCORE_HEAD = ("순번", "학번", "이름", "총점", "석차")


def _choice(question: int) -> int:
    return question % 5 + 1


def _answers(printed: int, uncertain: tuple[int, ...] = ()) -> tuple[AnswerValue, ...]:
    """A sheet printing Q1..Q<printed>, every one marked; later questions read unasked."""
    return tuple(
        UNASKED
        if question > printed
        else AnswerValue(
            (_choice(question),),
            AnswerStatus.UNCERTAIN if question in uncertain else AnswerStatus.NORMAL,
        )
        for question in range(1, 101)
    )


def _imported(serial: int, answers: tuple[AnswerValue, ...]) -> ImportedResponseRef:
    return ImportedResponseRef(
        1,
        f"work-{serial}",
        "a" * 64,
        "recognition",
        serial + 1,
        serial - 1,
        serial,
        f"scan-{serial}.png",
        f"2026000{serial}",
        "미등록",
        answers,
        "",
    )


def _effective(serial: int, answers: tuple[AnswerValue, ...]) -> EffectiveResponse:
    return EffectiveResponse(
        f"work-{serial}",
        SourceKind.IMAGE,
        f"scan-{serial}.png",
        f"2026000{serial}",
        StudentIdStatus.NORMAL,
        answers,
    )


def _key(asked: int) -> AnswerKeySnapshot:
    entries = tuple(
        AnswerKeyEntry(
            question,
            AnswerValue((_choice(question),), AnswerStatus.NORMAL),
            "1",
            KeyQuestionStatus.ANSWER,
        )
        if question <= asked
        else AnswerKeyEntry(question, UNASKED, "0", KeyQuestionStatus.UNASKED)
        for question in range(1, 101)
    )
    return AnswerKeySnapshot(
        1, AnswerKeySnapshotKind.WORKBOOK, "key.xlsx", "a" * 64, "정답표", "v1", entries, ()
    )


def _parse(path: Path) -> Ok | Err:
    with path.open("rb") as source:
        return parse_response_book(
            source,
            sheet_name=RESPONSE_SHEET_NAME,
            session_id="import-session",
            source_sha256=sha256(path.read_bytes()).hexdigest(),
        )


def _painted_columns(sheet) -> set[int]:
    return {
        cell.column
        for row in sheet.iter_rows(min_row=2)
        for cell in row
        if cell.fill.fill_type is not None
    }


def test_the_books_show_questions_up_to_the_last_one_a_sheet_prints() -> None:
    assert question_count([_answers(40), _answers(40)]) == 40
    assert question_count([_answers(40), _answers(60)]) == 60
    assert question_count([_answers(100)]) == 100
    assert question_count([]) == 100
    assert response_headers(40) == (*RESPONSE_FIXED_HEADERS, *question_headers(40), "비고")


@pytest.mark.parametrize("printed", (40, 60, 100))
def test_the_response_book_has_one_column_per_question_and_imports_again(
    tmp_path: Path, printed: int
) -> None:
    rows = (_imported(1, _answers(printed)), _imported(2, _answers(printed)))
    path = tmp_path / "responses.xlsx"

    write_response_projection(path, rows, session_id="s", revision=1, manifest_sha256="0" * 64)

    sheet = load_workbook(path)[RESPONSE_SHEET_NAME]
    assert tuple(cell.value for cell in sheet[1]) == response_headers(printed)
    assert sheet.max_column == printed + 5
    parsed = _parse(path)
    assert isinstance(parsed, Ok)
    for row in parsed.value:
        assert len(row.answers) == 100
        assert [answer.choices for answer in row.answers[:printed]] == [
            (_choice(question),) for question in range(1, printed + 1)
        ]
        assert all(answer.status is AnswerStatus.NORMAL for answer in row.answers[:printed])
        assert all(answer.status is AnswerStatus.UNASKED for answer in row.answers[printed:])


@pytest.mark.parametrize(
    "questions",
    (
        (),
        ("Q1", "Q3"),
        ("Q2",),
        ("Q1", "Q2", "비고"),
        tuple(f"Q{number}" for number in range(1, 102)),
    ),
    ids=("none", "gap", "not-from-1", "note-not-last", "over-100"),
)
def test_a_book_whose_question_columns_do_not_run_1_to_n_is_refused(
    tmp_path: Path, questions: tuple[str, ...]
) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = RESPONSE_SHEET_NAME
    headers = [*RESPONSE_FIXED_HEADERS, *questions, "비고"]
    if questions == ("Q1", "Q2", "비고"):
        headers = [*RESPONSE_FIXED_HEADERS, "Q1", "Q2", "비고", "메모"]
    sheet.append(headers)
    sheet.append([1, "scan.png", "20260001", "", *("1" for _ in headers[4:-1]), ""])
    path = tmp_path / "book.xlsx"
    workbook.save(path)

    parsed = _parse(path)

    assert isinstance(parsed, Err)
    assert parsed.errors[0].code == "XLSX_HEADER_INVALID"


def test_the_score_books_stop_at_the_last_question_of_the_form(tmp_path: Path) -> None:
    wrong_last = list(_answers(60))
    wrong_last[59] = AnswerValue((_choice(60) % 5 + 1,), AnswerStatus.NORMAL)
    responses = (
        _effective(1, _answers(60, uncertain=(7,))),
        _effective(2, tuple(wrong_last)),
    )
    key = _key(60)
    common = dict(
        exam_name="sixty",
        committed_at="2026-10-20T01:02:03.123456Z",
        session_id="session-60",
        revision=2,
        manifest_sha256="b" * 64,
        responses=responses,
        key=key,
        scores=score_effective(ScoreInput(responses, key)),
    )

    score = load_workbook(write_score_book(tmp_path, **common))
    final = load_workbook(write_final_book(tmp_path, **common))

    headers = (*SCORE_HEAD, *question_headers(60), "비고")
    for title in ("채점결과", "결과OX"):
        sheet = score[title]
        assert tuple(cell.value for cell in sheet[1]) == headers
        assert sheet.max_column == 66
        # Q7 of the first student awaits a check (Q7 and the total yellow), the second got Q60
        # wrong (pink); nothing past Q60, so 비고 in column 66 stays plain.
        assert _painted_columns(sheet) == {4, 12, 65}
    assert tuple(cell.value for cell in final["최종성적표"][1]) == (
        *headers,
        "수정여부",
        "수정문항",
        "확정일시",
    )
    assert final["최종성적표"].max_column == 69
    assert tuple(cell.value for cell in score[RESPONSE_SHEET_NAME][1]) == response_headers(60)
    assert score[RESPONSE_SHEET_NAME].max_column == 65
    evidence = score["판독 근거"]
    assert tuple(cell.value for cell in evidence[1])[:63] == (
        "순번",
        "학번",
        "이름",
        *question_headers(60),
    )
    assert evidence.max_column == 64
    assert score["정답표"].max_row == 61
    assert score["문항 분석"].max_row == 61


def test_the_response_tab_of_a_short_score_book_starts_a_new_exam(tmp_path: Path) -> None:
    responses = (_effective(1, _answers(40)), _effective(2, _answers(40)))
    key = _key(40)
    path = write_score_book(
        tmp_path,
        exam_name="forty",
        committed_at="2026-10-20T01:02:03.123456Z",
        session_id="session-40",
        revision=2,
        manifest_sha256="b" * 64,
        responses=responses,
        key=key,
        scores=score_effective(ScoreInput(responses, key)),
    )

    parsed = _parse(path)

    assert isinstance(parsed, Ok)
    assert [row.answers for row in parsed.value] == [_answers(40), _answers(40)]


def test_an_old_full_width_book_reads_its_empty_tail_as_blank(tmp_path: Path) -> None:
    # Books from before 4.4.0 always have Q1..Q100; empty answer cells there are blank, so
    # such a book counts as a 100-question form (a known limit, see docs/releases/4.4.0.md).
    rows = (_imported(1, _answers(100)),)
    path = tmp_path / "old.xlsx"
    write_response_projection(path, rows, session_id="s", revision=1, manifest_sha256="0" * 64)
    book = load_workbook(path)
    sheet = book[RESPONSE_SHEET_NAME]
    for column in range(45, 105):
        sheet.cell(2, column).value = None
    book.save(path)

    parsed = _parse(path)

    assert isinstance(parsed, Ok)
    (row,) = parsed.value
    assert all(answer.status is AnswerStatus.BLANK for answer in row.answers[40:])


def test_the_score_book_keeps_every_question_the_key_asks() -> None:
    responses = (_effective(1, _answers(40)),)

    assert book_question_count(responses, _key(40)) == 40
    assert book_question_count(responses, _key(45)) == 45
    assert book_question_count(responses, _key(30)) == 30
    assert book_question_count((_effective(1, _answers(100)),), _key(50)) == 50


def test_the_key_summary_counts_against_the_form() -> None:
    printed = tuple(range(1, 41))

    exact = GradingPresenter.validation(AnswerKeyValidation(_key(40)), printed)
    fewer = GradingPresenter.validation(AnswerKeyValidation(_key(35)), printed)
    longer = GradingPresenter.validation(AnswerKeyValidation(_key(50)), printed)
    unknown_form = GradingPresenter.validation(AnswerKeyValidation(_key(40)))

    assert (exact.question_count, exact.unasked_count, exact.errors) == (40, 0, ())
    # Rows after the key's last question are simply not on the exam, not 미출제.
    assert (fewer.question_count, fewer.unasked_count, fewer.errors) == (35, 0, ())
    # The summary shows what this form would grade; the error names the rows past it.
    assert (longer.question_count, longer.unasked_count, longer.total_points) == (40, 0, "40")
    assert longer.errors == (
        "정답표에 41~50번이 있지만 답안지에는 40번까지만 있습니다. "
        "정답표에서 41~50번 행을 지운 뒤 다시 불러오세요.",
    )
    assert (unknown_form.question_count, unknown_form.unasked_count) == (40, 0)


def test_the_key_summary_on_a_longer_sheet_counts_only_the_exam() -> None:
    # A 50-question exam on the 100-question sheet: 50 asked, nothing 미출제.
    hundred = tuple(range(1, 101))
    fifty = GradingPresenter.validation(AnswerKeyValidation(_key(50)), hundred)
    # A blank row inside the key is 미출제; it is still part of the exam's range.
    gap = _key(40)
    entries = list(gap.entries)
    entries[29] = AnswerKeyEntry(30, UNASKED, "0", KeyQuestionStatus.UNASKED)
    with_gap = GradingPresenter.validation(
        AnswerKeyValidation(replace(gap, entries=tuple(entries))), hundred
    )

    assert (fifty.question_count, fifty.unasked_count, fifty.total_points) == (50, 0, "50")
    assert fifty.errors == ()
    assert (with_gap.question_count, with_gap.unasked_count, with_gap.total_points) == (
        39,
        1,
        "39",
    )


def test_a_fifty_question_key_on_a_hundred_question_sheet_gives_fifty_columns(
    tmp_path: Path,
) -> None:
    # Everyone leaves Q51..Q100 blank; one student has an unconfirmed mark at Q57.
    blank = AnswerValue((), AnswerStatus.BLANK)
    answers = tuple(
        answer if question <= 50 else blank
        for question, answer in enumerate(_answers(100), 1)
    )
    stray = list(answers)
    stray[56] = AnswerValue((2,), AnswerStatus.UNCERTAIN)
    responses = (_effective(1, answers), _effective(2, tuple(stray)))
    key = _key(50)

    path = write_score_book(
        tmp_path,
        exam_name="fifty",
        committed_at="2026-10-21T01:02:03.123456Z",
        session_id="session-50",
        revision=2,
        manifest_sha256="b" * 64,
        responses=responses,
        key=key,
        scores=score_effective(ScoreInput(responses, key)),
    )

    book = load_workbook(path)
    for title in ("채점결과", "결과OX"):
        assert book[title].max_column == 56
        assert book[title].cell(1, 55).value == "Q50"
        assert _painted_columns(book[title]) == set()
    responses_sheet = book[RESPONSE_SHEET_NAME]
    assert tuple(cell.value for cell in responses_sheet[1]) == response_headers(50)
    # Q57 does not count, so it is neither shown nor named; the book imports as it is.
    assert [responses_sheet.cell(row, 55).value for row in (2, 3)] == [None, None]
    assert book["판독 근거"].max_column == 54
    assert book["정답표"].max_row == 51
    assert book["채점결과"].cell(2, 4).value == 50
    assert isinstance(_parse(path), Ok)


def test_the_grading_use_case_reports_the_printed_questions() -> None:
    snapshot = SimpleNamespace(responses=(_effective(1, _answers(40)), _effective(2, _answers(40))))

    class Snapshots:
        def read_grading_snapshot(self, session_id: str, revision: int) -> Ok:
            return Ok(snapshot)

    use_case = GradingUseCase(Snapshots(), None, None)  # type: ignore[arg-type]

    assert use_case.form_questions("session-40", 1) == Ok(tuple(range(1, 41)))
