from dataclasses import replace
from decimal import Decimal
from zipfile import ZipFile

import pytest
from openpyxl import load_workbook
from tests.integration.test_score_workbook import (
    _marked,
    _scoring_key,
    _student,
    _write,
)
from tests.unit.test_response_book import YELLOW, _painted, _uncertain

from omr_grader.application.dto import ScoreInput
from omr_grader.domain.enums import CellStatus, StudentIdStatus
from omr_grader.domain.grading import score_effective
from omr_grader.domain.models import CellEvidence
from omr_grader.workbooks.answer_key import answer_key_rows
from omr_grader.workbooks.schemas import RESPONSE_SHEET_NAME
from omr_grader.workbooks.score_book import (
    answer_margins,
    scored_image_names,
    student_order,
    write_final_book,
    write_score_book,
)

_COMMITTED = "2026-07-28T01:02:03.123456Z"


def _classroom():
    """Twelve scored students and one with an unconfirmed Q8.

    Student k answers Q(6+j) correctly (a 1) exactly when k > j, otherwise marks a 2, so the
    totals rise with k.  Student 0 leaves Q7 blank and student 1 marks 1 and 2 on it.
    """
    students = []
    for k in range(12):
        changes = {6 + j: _marked(2) for j in range(12) if k <= j}
        if k == 0:
            changes[7] = _marked()
        if k == 1:
            changes[7] = _marked(1, 2)
        students.append(_student(f"wi_{k:02d}", f"2024{k:04d}", changes))
    students.append(_student("wi_review", "20240099", {8: _uncertain(2)}))
    return students


def _items(book):
    sheet = book["문항 분석"]
    return sheet, {row[0].value: row for row in sheet.iter_rows(min_row=2)}


def _values(row):
    return [cell.value for cell in row]


def test_item_analysis_counts_choices_rates_and_discrimination(tmp_path):
    book = _write(write_score_book, tmp_path, tuple(_classroom()))
    sheet, rows = _items(book)

    assert [cell.value for cell in sheet[1]] == [
        "문항",
        "정답",
        "배점",
        "응시자 수",
        "정답률(%)",
        "①",
        "②",
        "③",
        "④",
        "⑤",
        "무응답",
        "중복",
        "확인 필요",
        "변별도",
    ]
    # Q1..Q20 are asked; Q21..Q100 are not.
    assert sorted(rows) == list(range(1, 21))
    assert sheet.freeze_panes == "A2" and all(cell.font.b for cell in sheet[1])
    # Q6: 12 of 13 right.  Upper group (k=9..11) 3/3, lower group (k=0..2) 2/3.
    assert _values(rows[6]) == [6, 1, 1, 13, 92.3, 12, 1, 0, 0, 0, 0, 0, 0, 0.33]
    # Q7: student 0 left it blank, student 1 marked two; the others are right (lower group 1/3).
    assert _values(rows[7]) == [7, 1, 1, 13, 84.6, 11, 0, 0, 0, 0, 1, 1, 0, 0.67]
    # Q8: the unconfirmed answer is not counted among the takers.
    assert _values(rows[8])[3:5] == [12, 75.0] and rows[8][12].value == 1
    assert rows[11][13].value == 1.0
    # Q17 was right for nobody who has a total, and the reviewer alone got it.
    assert _values(rows[17]) == [17, 1, 1, 13, 7.7, 1, 12, 0, 0, 0, 0, 0, 0, 0.0]
    # Q5 wants 1 and 2 together: shown as text, and the right answers are no 중복.
    assert _values(rows[5])[:5] == [5, "1,2", 1, 13, 100.0]
    assert rows[5][11].value == 0
    # Q4's key is a single number.
    assert rows[4][1].value == 3 and rows[4][1].data_type == "n"
    assert rows[4][13].number_format == "0.00" and rows[6][4].number_format == "0.0"


def test_item_analysis_highlights_hard_questions_strong_distractors_and_flat_discrimination(
    tmp_path,
):
    _, rows = _items(_write(write_score_book, tmp_path, tuple(_classroom())))

    # Q17: under 30% correct, choice 2 beaten the answer, discrimination 0.
    assert _painted_columns(rows[17]) == {5: YELLOW, 7: YELLOW, 14: YELLOW}
    # Q16: 2 of 13 right, choice 2 picked by 11.
    assert _painted_columns(rows[16]) == {5: YELLOW, 7: YELLOW}
    # Q11 is half right with the wrong choice at 6 against 7 right.
    assert _painted_columns(rows[11]) == {}
    assert _painted_columns(rows[6]) == {}
    # Everyone answered Q1 right: discrimination is 0 but that separates nobody, so no flag ...
    assert rows[1][13].value == 0.0 and _painted_columns(rows[1]) == {}
    # ... but the everyone-correct Q3 stays plain with 100% correct.
    assert _values(rows[3])[:5] == [3, "전원", 1, 13, 100.0]
    assert _painted_columns(rows[3]) == {}


def _painted_columns(row):
    return {
        cell.column: (cell.fill.fgColor.rgb, cell.font.color.rgb)
        for cell in row
        if cell.fill.fill_type is not None
    }


def test_discrimination_is_blank_below_ten_scored_students(tmp_path):
    students = tuple(_classroom()[2:11]) + (_classroom()[-1],)  # nine scored and one reviewer
    sheet, rows = _items(_write(write_score_book, tmp_path, students))

    assert all(row[13].value is None for row in rows.values())
    assert rows[6][3].value == 10
    assert not any(row[13].fill.fill_type is not None for row in sheet.iter_rows(min_row=2))


def test_final_book_carries_the_same_analysis_sheet_after_the_outcome_sheet(tmp_path):
    book = _write(write_final_book, tmp_path, tuple(_classroom()))

    assert book.sheetnames == [
        "최종성적표",
        "결과OX",
        "문항 분석",
        RESPONSE_SHEET_NAME,
        "판독 근거",
        "정답표",
        "색 설명",
    ]
    assert book["문항 분석"].max_row == 21


def test_book_order_names_and_image_names_share_one_rule():
    students = (
        _student("wi_c", "20240003", {}),
        _student("wi_unnamed", "20240001", {}),
        _student("wi_a", "20240009", {}),
        _student("wi_noid", "20240000", {}),
        _student("wi_slash", "20240005", {}),
    )
    students = tuple(
        replace(item, student_id=None, student_id_status=StudentIdStatus.UNREADABLE)
        if item.work_item_id == "wi_noid"
        else item
        for item in students
    )
    names = {"20240003": "다현", "20240009": "가 윤", "20240005": "김:나/ㅅ"}

    ordered = student_order(students, names)

    assert [item.work_item_id for item in ordered] == [
        "wi_a",
        "wi_slash",
        "wi_c",
        "wi_unnamed",
        "wi_noid",
    ]
    assert scored_image_names(students, names) == {
        "wi_a": "001_20240009_가_윤.jpg",
        "wi_slash": "002_20240005_김_나_ㅅ.jpg",
        "wi_c": "003_20240003_다현.jpg",
        "wi_unnamed": "004_20240001.jpg",
        "wi_noid": "005_학번확인필요.jpg",
    }


def test_score_sheet_rows_follow_student_order(tmp_path):
    students = (
        _student("wi_c", "20240003", {}),
        _student("wi_a", "20240009", {}),
        _student("wi_b", "20240002", {}),
    )
    names = {"20240003": "다현", "20240009": "가윤"}
    key = _scoring_key()
    path = write_score_book(
        tmp_path,
        exam_name="시험",
        committed_at=_COMMITTED,
        session_id="s",
        revision=1,
        manifest_sha256="b" * 64,
        responses=students,
        key=key,
        scores=score_effective(ScoreInput(students, key)),
        names_by_student_id=names,
    )
    sheet = load_workbook(path)["채점결과"]

    ordered = student_order(students, names)
    assert [sheet.cell(row, 2).value for row in (2, 3, 4)] == [item.student_id for item in ordered]


def test_single_answers_are_numbers_and_several_stay_text(tmp_path):
    book = _write(write_score_book, tmp_path, tuple(_classroom()))
    sheet = book["채점결과"]

    # Q1 of the first student is a number; the ID stays text.
    assert sheet.cell(2, 6).value == 1 and sheet.cell(2, 6).data_type == "n"
    assert sheet.cell(2, 2).value == "20240000" and sheet.cell(2, 2).data_type == "s"
    # Q5 wants two choices; student 0 left Q7 blank and student 1 marked two choices on it.
    assert (sheet.cell(2, 10).value, sheet.cell(2, 10).data_type) == ("1,2", "s")
    assert [sheet.cell(2, 12).value, sheet.cell(3, 12).value] == [None, "1,2"]
    # The answer key sheet follows the same rule; the 응답원본 sheet stays text for import.
    key_sheet = book["정답표"]
    assert [key_sheet.cell(number, 2).value for number in (2, 4, 5, 6, 7, 22)] == [
        1,
        "전원",
        3,
        "1,2",
        1,
        None,
    ]
    assert key_sheet.cell(2, 2).data_type == "n"
    assert book[RESPONSE_SHEET_NAME].cell(2, 5).value == "1"
    assert book[RESPONSE_SHEET_NAME].cell(2, 5).data_type == "s"


def test_the_stored_answer_key_copy_still_spells_everyone_correct_as_zero():
    rows = answer_key_rows(_scoring_key())

    assert rows[2] == (3, "0", Decimal("1"))
    assert rows[3] == (4, "3", Decimal("1"))


@pytest.mark.parametrize("writer", (write_score_book, write_final_book))
def test_text_columns_are_exempt_from_the_number_stored_as_text_warning(tmp_path, writer):
    key = _scoring_key()
    students = tuple(_classroom())
    path = writer(
        tmp_path,
        exam_name="시험",
        committed_at=_COMMITTED,
        session_id="s",
        revision=1,
        manifest_sha256="b" * 64,
        responses=students,
        key=key,
        scores=score_effective(ScoreInput(students, key)),
    )
    book = load_workbook(path)
    with ZipFile(path) as package:
        sheets = {
            title: package.read(f"xl/worksheets/sheet{index}.xml").decode()
            for index, title in enumerate(book.sheetnames, 1)
        }

    score_title = book.sheetnames[0]
    assert 'sqref="B2:B14 F2:DA14" numberStoredAsText="1"' in sheets[score_title]
    assert 'sqref="B2:B14" numberStoredAsText="1"' in sheets["결과OX"]
    assert 'sqref="C2:C14 E2:CZ14" numberStoredAsText="1"' in sheets[RESPONSE_SHEET_NAME]
    assert 'sqref="B2:B14" numberStoredAsText="1"' in sheets["판독 근거"]
    assert 'sqref="B2:B101" numberStoredAsText="1"' in sheets["정답표"]
    assert 'sqref="B2:B21" numberStoredAsText="1"' in sheets["문항 분석"]
    assert "ignoredErrors" not in sheets["색 설명"]
    assert all(
        xml.index("<ignoredErrors>") > xml.index("</sheetData>")
        for title, xml in sheets.items()
        if "ignoredErrors" in xml
    )


def _evidence(question, fills):
    fills = [format(Decimal(fill).normalize(), "f") for fill in fills]
    return [
        CellEvidence(
            index=0,
            question=question,
            digit=None,
            choice=choice,
            pixel_rect=None,
            ratio_rect=None,
            fill_score=fill,
            selected=False,
            status=CellStatus.NORMAL,
        )
        for choice, fill in enumerate(fills, 1)
    ]


def test_answer_margins_are_the_lead_of_the_darkest_choice():
    cells = [
        *_evidence(1, ["0.70", "0.05", "0.02", "0.00", "0.01"]),
        *_evidence(2, ["0.40", "0.35", "0.02", "0.00", "0.01"]),
        *_evidence(3, ["0.02", "0.01"]),
        *_evidence(4, ["0.50"]),
    ]

    margins = answer_margins(cells)

    assert margins[1] == pytest.approx((0.70, 0.65))
    assert margins[2] == pytest.approx((0.40, 0.05))
    assert margins[3] == pytest.approx((0.02, 0.01))
    assert 4 not in margins


def test_evidence_sheet_shows_margins_and_leaves_imported_rows_blank(tmp_path):
    scanned = _student("wi_scan", "20240001", {})
    imported = _student("wi_import", "20240002", {})
    evidence = {
        "wi_scan": [
            *_evidence(1, ["0.70", "0.05", "0.02", "0.00", "0.01"]),
            *_evidence(2, ["0.40", "0.35", "0.02", "0.00", "0.01"]),
            # Nothing marked: a tiny margin is normal and stays plain.
            *_evidence(3, ["0.03", "0.02", "0.01", "0.00", "0.00"]),
            *_evidence(4, ["0.60", "0.50", "0.02", "0.00", "0.01"]),
        ]
    }
    key = _scoring_key()
    path = write_score_book(
        tmp_path,
        exam_name="시험",
        committed_at=_COMMITTED,
        session_id="s",
        revision=1,
        manifest_sha256="b" * 64,
        responses=(scanned, imported),
        key=key,
        scores=score_effective(ScoreInput((scanned, imported), key)),
        evidence_by_work_item=evidence,
    )
    sheet = load_workbook(path)["판독 근거"]

    headers = [cell.value for cell in sheet[1]]
    assert headers[:5] == ["순번", "학번", "이름", "Q1", "Q2"]
    assert headers[102] == "Q100"
    assert headers[103] == "가장 진한 칸과 두 번째 칸의 차이, 작을수록 판독이 애매함"
    assert all(cell.font.b for cell in sheet[1]) and sheet.freeze_panes == "D2"
    assert [sheet.cell(2, column).value for column in (1, 2, 4, 5, 6, 7, 8)] == [
        1,
        "20240001",
        0.65,
        0.05,
        0.01,
        0.1,
        None,
    ]
    assert sheet.cell(2, 4).number_format == "0.00"
    # Q2 (0.05) is yellow; Q4's lead of 0.10 is not under the limit; Q3 is unmarked.
    assert _painted(sheet, 2) == {5: YELLOW}
    # The imported student has no recognition evidence.
    assert [sheet.cell(3, column).value for column in (1, 2, 4, 5)] == [2, "20240002", None, None]
    assert _painted(sheet, 3) == {}


def test_legend_explains_the_new_sheets(tmp_path):
    legend = _write(write_score_book, tmp_path, tuple(_classroom()))["색 설명"]
    text = "\n".join(str(cell.value) for row in legend.iter_rows() for cell in row)

    for phrase in (
        "정답률 30% 미만",
        "변별도 0 이하",
        "상위 27%",
        "가장 진한 칸과 두 번째 칸의 차이, 작을수록 판독이 애매함",
        "정답표의 '전원' = 모든 학생 정답 처리",
    ):
        assert phrase in text
