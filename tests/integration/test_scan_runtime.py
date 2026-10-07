from __future__ import annotations

from dataclasses import replace
from hashlib import sha256
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook
from tests.unit.test_domain_models import answer, id_cell, page_ref

from omr_grader.application.dto import (
    CancelOperationCommand,
    ScanCommand,
    ScanProgress,
    ScanSource,
)
from omr_grader.application.scan_use_case import ScanUseCase
from omr_grader.domain.enums import (
    AnswerStatus,
    ExamTerm,
    ProcessingStatus,
    RosterSnapshotKind,
    StudentIdStatus,
)
from omr_grader.domain.errors import Err, Ok
from omr_grader.domain.models import (
    AnswerValue,
    AutomaticPage,
    RosterSnapshot,
    StudentIdRecognition,
)
from omr_grader.infrastructure.scan_runtime import ScanRuntime
from omr_grader.recognition.pipeline import PipelineSuccess, RecognitionArtifacts
from omr_grader.ui.workers import WorkerBatchResult, WorkerResult
from omr_grader.workbooks.response_import import parse_response_book
from omr_grader.workbooks.schemas import RESPONSE_SHEET_NAME
from tests.helpers.omr_engine import reference_sheet
from tests.helpers.synthetic_omr import write_png


def _command(source: Path, *, profile: str = "profile.omrtemplate") -> ScanCommand:
    return ScanCommand(
        "scan-session",
        "scan-operation",
        0,
        "Math",
        2026,
        ExamTerm.FIRST,
        profile,
        None,
        ScanSource((str(source),)),
        5,
        False,
    )


def test_scan_runtime_rejects_profile_source_with_wrong_extension_before_ingestion(
    tmp_path: Path,
) -> None:
    source = tmp_path / "scan.png"
    source.write_bytes(b"not inspected")

    class Profiles:
        def load(self, filename: str):
            return Ok(object())

    runtime = ScanRuntime(Profiles(), object())
    result = runtime.build_tasks(_command(source, profile="profile.json"))

    assert isinstance(result, Err)
    assert result.errors[0].code == "PROFILE_SOURCE_INVALID"


def test_scan_use_case_commits_only_complete_uncancelled_worker_results(tmp_path: Path) -> None:
    command = _command(tmp_path / "unused.png")
    tasks = (object(), object())

    class Source:
        def build_tasks(self, received, progress=None):
            assert received is command
            return Ok(tasks)

    class Worker:
        def run(self, received, *, multiprocessing, progress):
            assert received == tasks
            return WorkerBatchResult((), False)

    class Coordinator:
        def commit_scan(self, command, results):
            raise AssertionError("incomplete worker output must not publish")

    result = ScanUseCase(Source(), Worker).run_scan(command, Coordinator())

    assert isinstance(result, Err)
    assert result.errors[0].code == "WORKER_RESULT_INCOMPLETE"


def test_scan_use_case_reports_preparing_then_reading_then_saving(tmp_path: Path) -> None:
    command = _command(tmp_path / "unused.png")
    tasks = (object(), object())
    reports: list[ScanProgress] = []

    class Source:
        def build_tasks(self, received, progress=None):
            progress(1, 2)
            progress(2, 2)
            return Ok(tasks)

    class Worker:
        def run(self, received, *, multiprocessing, progress):
            progress(ScanProgress(2, 2, 0, 5, 0))
            return WorkerBatchResult(
                (WorkerResult(0, "first"), WorkerResult(1, "second")), False
            )

    class Coordinator:
        def commit_scan(self, command, results):
            assert reports[-1].phase == "save"
            return Ok(results)

    result = ScanUseCase(Source(), Worker).run_scan(
        command, Coordinator(), progress=reports.append
    )

    assert result == Ok(("first", "second"))
    assert [(item.phase, item.completed, item.total) for item in reports] == [
        ("prepare", 1, 2),
        ("prepare", 2, 2),
        ("recognize", 2, 2),
        ("save", 2, 2),
    ]


def test_cancelling_while_pages_are_prepared_stops_before_recognition(tmp_path: Path) -> None:
    command = _command(tmp_path / "unused.png")

    class Source:
        def build_tasks(self, received, progress=None):
            progress(1, 3)
            assert use_case.cancel_scan(CancelOperationCommand(received.operation_id)) == Ok(
                None
            )
            progress(2, 3)
            raise AssertionError("preparation must stop once the scan is cancelled")

    class Worker:
        def run(self, received, *, multiprocessing, progress):
            raise AssertionError("a cancelled scan must not be recognised")

        def cancel(self) -> None:
            pass

    class Coordinator:
        def commit_scan(self, command, results):
            raise AssertionError("a cancelled scan must not publish")

    use_case = ScanUseCase(Source(), Worker)
    result = use_case.run_scan(command, Coordinator())

    assert isinstance(result, Err)
    assert result.errors[0].code == "OPERATION_CANCELLED"
    # The operation is released, so a late cancel finds nothing to stop.
    cancelled = use_case.cancel_scan(CancelOperationCommand(command.operation_id))
    assert isinstance(cancelled, Err)
    assert cancelled.errors[0].code == "OPERATION_NOT_FOUND"


def test_a_cancel_after_recognition_stops_before_the_results_are_saved(
    tmp_path: Path,
) -> None:
    command = _command(tmp_path / "unused.png")

    class Source:
        def build_tasks(self, received, progress=None):
            return Ok((object(),))

    class Worker:
        def run(self, received, *, multiprocessing, progress):
            # The cancel arrives as the last page finishes.
            use_case.cancel_scan(CancelOperationCommand(command.operation_id))
            return WorkerBatchResult((WorkerResult(0, "page"),), False)

        def cancel(self) -> None:
            pass

    class Coordinator:
        def commit_scan(self, command, results):
            raise AssertionError("a cancelled scan must not publish")

    use_case = ScanUseCase(Source(), Worker)
    result = use_case.run_scan(command, Coordinator())

    assert isinstance(result, Err)
    assert result.errors[0].code == "OPERATION_CANCELLED"


def test_scan_runtime_reports_each_prepared_page(tmp_path: Path) -> None:
    sheet = reference_sheet()
    folder = tmp_path / "scans"
    folder.mkdir()
    for name in ("a.png", "b.png", "c.png"):
        write_png(folder / name, sheet.image)
    reports: list[tuple[int, int]] = []

    class Profiles:
        def load(self, filename: str):
            return Ok(sheet.profile)

    result = ScanRuntime(Profiles(), object()).build_tasks(
        _command(folder), lambda done, total: reports.append((done, total))
    )

    assert isinstance(result, Ok)
    assert len(result.value) == 3
    assert reports == [(1, 3), (2, 3), (3, 3)]


def test_scan_artifacts_preserve_original_pdf_bytes(tmp_path: Path) -> None:
    source = tmp_path / "source scan.pdf"
    payload = b"%PDF-1.7 exact source bytes"
    source.write_bytes(payload)
    runtime = ScanRuntime(object(), object())
    roster = RosterSnapshot(
        1,
        RosterSnapshotKind.NONE,
        None,
        None,
        None,
        "v1",
        (),
        (),
    )

    result = runtime._artifacts(
        _command(source),
        "a" * 64,
        roster,
        (),
        "2026-07-31T00:00:00.000000Z",
    )

    assert isinstance(result, Ok)
    artifacts = result.value[0]
    assert artifacts["sources/scans/001_source_scan.pdf"] == payload


def _recognized(
    number: int, *, uncertain: tuple[int, ...] = (), student_id: str | None = "20240001"
) -> PipelineSuccess:
    """One recognized page; ``student_id=None`` stands for an ID the reader could not use."""
    if student_id is None:
        identity = StudentIdRecognition(
            None, StudentIdStatus.INVALID, tuple(id_cell(None) for _ in range(8))
        )
    else:
        identity = StudentIdRecognition(
            student_id, StudentIdStatus.NORMAL, tuple(id_cell(digit) for digit in student_id)
        )
    unconfirmed = AnswerValue((1,), AnswerStatus.UNCERTAIN)
    answers = tuple(
        answer(question, unconfirmed if question in uncertain else None)
        for question in range(1, 101)
    )
    evidence = tuple(cell for item in identity.cells for cell in item.candidates) + tuple(
        cell for item in answers for cell in item.cells
    )
    reference = replace(
        page_ref(),
        work_item_id=f"work-{number}",
        source_display_name=f"scan-{number}.png",
        input_ordinal=number - 1,
        artifact_stem=f"scan-{number}",
    )
    page = AutomaticPage(
        1,
        reference,
        ProcessingStatus.NEEDS_MANUAL_REVIEW,
        0,
        "0.99",
        "0.98",
        (1000, 1400),
        ("1",) * 9,
        ("1",) * 9,
        identity,
        answers,
        evidence,
    )
    return PipelineSuccess(page, RecognitionArtifacts(b"normalized", b"{}", b"overlay"))


def test_scan_response_book_notes_unconfirmed_answers_and_unusable_student_ids(
    tmp_path: Path,
) -> None:
    roster = RosterSnapshot(1, RosterSnapshotKind.NONE, None, None, None, "v1", (), ())
    results = (
        _recognized(1, uncertain=(3, 17)),
        _recognized(2, student_id=None),
        _recognized(3, uncertain=(3, 4, 5), student_id=None),
        _recognized(4),
    )

    result = ScanRuntime(object(), object())._artifacts(
        _command(tmp_path / "scan.png"), "a" * 64, roster, results, "2026-07-31T00:00:00.000000Z"
    )

    assert isinstance(result, Ok)
    workbook = result.value[0]["responses.xlsx"]
    sheet = load_workbook(BytesIO(workbook))[RESPONSE_SHEET_NAME]
    assert [
        (sheet.cell(row, 2).value, sheet.cell(row, 3).value, sheet.cell(row, 105).value)
        for row in range(2, 6)
    ] == [
        ("scan-1.png", "20240001", "확인 필요: 3, 17번"),
        ("scan-2.png", None, "학번 확인 필요"),
        ("scan-3.png", None, "확인 필요: 3~5번 / 학번 확인 필요"),
        ("scan-4.png", "20240001", None),
    ]
    # Unconfirmed answers hold the engine's guess, so the book is not imported as is.
    refused = parse_response_book(
        BytesIO(workbook),
        sheet_name=RESPONSE_SHEET_NAME,
        session_id="import-session",
        source_sha256=sha256(workbook).hexdigest(),
    )
    assert isinstance(refused, Err)
    assert refused.errors[0].code == "XLSX_REVIEW_PENDING"
    # Yellow marks exactly the unconfirmed answers: Q3 and Q17 of page 1, Q3..Q5 of page 3.
    yellow = {
        (cell.row, cell.column)
        for row in sheet.iter_rows(min_row=2, min_col=5, max_col=104)
        for cell in row
        if cell.fill.fill_type == "solid" and cell.fill.fgColor.rgb == "FFFFEB9C"
    }
    assert yellow == {(2, 7), (2, 21), (4, 7), (4, 8), (4, 9)}
