"""Two scanned and graded parts become one 합산 성적표 with a student in both and one in part 1."""

from __future__ import annotations

from pathlib import Path
from uuid import uuid4

import openpyxl

from omr_grader.application.answer_key_use_case import AnswerKeyWorkbookUseCase
from omr_grader.application.dto import RegradeCommand, ScanCommand, ScanSource
from omr_grader.application.grading_use_case import GradingUseCase
from omr_grader.bootstrap import bootstrap
from omr_grader.domain.enums import ExamTerm
from omr_grader.domain.errors import Err, Ok
from omr_grader.infrastructure.combined_report import (
    export_combined_report,
    load_part_scores,
    write_subject_config_sample,
)
from omr_grader.infrastructure.form_detection import FormDetector
from omr_grader.infrastructure.grading_runtime import CommittedGradingSnapshotReader
from omr_grader.infrastructure.paths import ManagedPaths
from omr_grader.infrastructure.profile_store import ProfileStore
from omr_grader.infrastructure.scan_runtime import ScanRuntime, bind_scan_runtime
from omr_grader.infrastructure.session_store import SessionCommitCoordinator, SessionStore
from omr_grader.workbooks.answer_key import ANSWER_KEY_HEADERS
from tests.helpers.pdf_writer import write_pdf
from tests.helpers.synthetic_omr import encode_png, render_sheet


def _answer(question: int) -> int:
    return (question * 7) % 5 + 1


def _right_up_to(count: int) -> dict[int, int]:
    """Correct for the first ``count`` questions, one choice off for the rest."""
    return {
        question: _answer(question) if question <= count else _answer(question) % 5 + 1
        for question in range(1, 101)
    }


class _Workshop:
    def __init__(self, tmp_path: Path) -> None:
        self.tmp_path = tmp_path
        (tmp_path / "root").mkdir()
        outcome = bootstrap(ManagedPaths.from_root(tmp_path / "root"))
        assert not isinstance(outcome, Err)
        self.paths, self.token = outcome.value.paths, outcome.value.capability_token
        self.profiles = ProfileStore(self.paths, self.token)
        self.store = SessionStore(self.paths)
        self.coordinator = SessionCommitCoordinator(self.store)
        self.runtime = bind_scan_runtime(
            ScanRuntime(ProfileStore(self.paths, self.token), self.store)
        )

    def _inputs(self, name: str, sheets: list[tuple[str, int]]) -> tuple[Path, Path]:
        folder = self.tmp_path / name
        folder.mkdir()
        pdf = folder / "scans.pdf"
        write_pdf(
            pdf,
            [
                encode_png(render_sheet(_right_up_to(right), student_id, seed=seed))
                for seed, (student_id, right) in enumerate(sheets, 1)
            ],
        )
        key = folder / "key.xlsx"
        book = openpyxl.Workbook()
        sheet = book.active
        assert sheet is not None
        sheet.title = "정답표"
        sheet.append(list(ANSWER_KEY_HEADERS))
        for question in range(1, 101):
            sheet.append([question, str(_answer(question)), 1])
        book.save(key)
        return pdf, key

    def graded_session(
        self, exam_name: str, sheets: list[tuple[str, int]], *, grade: bool = True
    ) -> str:
        pdf, key = self._inputs(exam_name, sheets)
        detected = FormDetector(self.profiles).detect((str(pdf),))
        assert isinstance(detected, Ok)
        profile = detected.value.profile_filename
        if detected.value.generated_profile is not None:
            saved = self.profiles.save_generated(
                detected.value.generated_profile, detected.value.suggested_filename
            )
            assert isinstance(saved, Ok)
            profile = saved.value.stored_name
        assert profile is not None
        scanned = self.runtime.run_scan(
            ScanCommand(
                f"scan-{uuid4().hex}",
                uuid4().hex,
                0,
                exam_name,
                None,
                ExamTerm.UNSPECIFIED,
                profile,
                None,
                ScanSource((str(pdf),)),
                5,
                False,
            )
        )
        assert isinstance(scanned, Ok)
        if not grade:
            return scanned.value.session_id
        graded = GradingUseCase(
            CommittedGradingSnapshotReader(self.coordinator),
            AnswerKeyWorkbookUseCase(),
            self.coordinator,
        ).regrade(
            RegradeCommand(
                scanned.value.session_id, scanned.value.revision, str(key), "정답표", uuid4().hex
            )
        )
        assert isinstance(graded, Ok)
        return graded.value.session_id


def _table(path: str, name: str = "합산결과") -> list[dict[str, object]]:
    rows = list(openpyxl.load_workbook(path)[name].iter_rows(values_only=True))
    return [dict(zip(rows[0], row, strict=True)) for row in rows[1:]]


def test_two_graded_parts_are_combined_with_a_student_missing_from_part_two(tmp_path):
    workshop = _Workshop(tmp_path)
    part1 = workshop.graded_session("졸업고사 1교시", [("20260001", 100), ("20260002", 60)])
    part2 = workshop.graded_session("졸업고사 2교시", [("20260001", 40)])

    loaded = load_part_scores(workshop.coordinator, part1, "파트1")
    assert isinstance(loaded, Ok)
    assert loaded.value.exam_name == "졸업고사 1교시"
    assert loaded.value.maximum == 100
    assert [row.total for row in loaded.value.rows] == [100, 60]
    assert [row.part_order for row in loaded.value.rows] == [1, 2]

    config = tmp_path / "subjects.xlsx"
    assert isinstance(write_subject_config_sample(str(config)), Ok)
    destination = tmp_path / "out" / "합산.xlsx"
    summary = export_combined_report(
        workshop.coordinator, (part1, part2), str(config), str(destination), "2026-02-01 10:00"
    )

    assert isinstance(summary, Ok)
    assert (summary.value.students, summary.value.complete) == (2, 1)
    assert (summary.value.needs_review, summary.value.unreadable) == (1, 0)
    assert summary.value.path == str(destination)
    rows = _table(str(destination))
    both = next(row for row in rows if row["학번"] == "20260001")
    only_first = next(row for row in rows if row["학번"] == "20260002")
    assert (both["파트1 점수"], both["파트2 점수"], both["총점"], both["통합 석차"]) == (
        100,
        40,
        140,
        1,
    )
    # The sample config: 해부생리학 = P1 1~30, 보건의료관계법규 = P1 91~100 + P2 1~10,
    # 작업치료평가 = P2 11~40, 아동작업치료 = P2 41~70, 정신사회작업치료 = P2 71~100.
    assert (both["해부생리학"], both["보건의료관계법규"], both["작업치료평가"]) == (30, 20, 30)
    assert (both["아동작업치료"], both["정신사회작업치료"]) == (0, 0)
    # 140 of 200 clears the total, but two subjects are below 60%.
    assert both["합격 여부"] == "불합격"
    assert only_first["총점"] is None and only_first["합격 여부"] is None
    assert only_first["비고"] == "확인 필요: 파트2 기록 없음 (학번 확인)"
    book = openpyxl.load_workbook(destination)
    sheet = book["합산결과"]
    assert sheet["A3"].fill.start_color.rgb == "FFFFEB9C"
    assert book.sheetnames[0] == "과목별 합격"
    subject_sheet = book["과목별 합격"]
    headers = [cell.value for cell in subject_sheet[1]]
    first, second = (
        dict(zip(headers, (cell.value for cell in subject_sheet[row]), strict=True))
        for row in (4, 5)
    )
    assert first["학번"] == "20260001" and first["판정"] == "불합격"
    assert first["미달 과목"] == "아동작업치료, 정신사회작업치료"
    assert second["판정"] == "확인 필요" and second["합계"] is None
    parts = _table(str(destination), "파트")
    assert [row["시험명"] for row in parts[:2]] == ["졸업고사 1교시", "졸업고사 2교시"]
    assert parts[0]["폴더"].endswith("졸업고사_1교시")
    assert [path.name for path in destination.parent.iterdir()] == ["합산.xlsx"]


def test_an_ungraded_exam_and_a_config_for_a_missing_part_are_refused(tmp_path):
    workshop = _Workshop(tmp_path)
    graded = workshop.graded_session("졸업고사 1교시", [("20260001", 100)])
    ungraded = workshop.graded_session("졸업고사 2교시", [("20260001", 100)], grade=False)
    destination = tmp_path / "합산.xlsx"

    refused = export_combined_report(
        workshop.coordinator, (graded, ungraded), None, str(destination), "now"
    )

    assert isinstance(refused, Err)
    assert refused.errors[0].context["reason"] == "채점하지 않은 시험이 있습니다: 졸업고사 2교시"
    assert not destination.exists()

    config = tmp_path / "subjects.xlsx"
    assert isinstance(write_subject_config_sample(str(config)), Ok)
    too_few = export_combined_report(
        workshop.coordinator, (graded,), str(config), str(destination), "now"
    )

    assert isinstance(too_few, Err)
    assert "파트2" in str(too_few.errors[0].context["reason"])
    assert not destination.exists()


def test_exporting_again_replaces_the_chosen_file_and_unknown_sessions_fail(tmp_path):
    workshop = _Workshop(tmp_path)
    graded = workshop.graded_session("졸업고사 1교시", [("20260001", 100)])
    destination = tmp_path / "합산.xlsx"
    destination.write_bytes(b"old")

    done = export_combined_report(workshop.coordinator, (graded,), None, str(destination), "now")

    assert isinstance(done, Ok) and done.value.complete == 1
    assert openpyxl.load_workbook(destination).sheetnames == ["합산결과", "파트"]
    missing = export_combined_report(
        workshop.coordinator, (graded, uuid4().hex), None, str(destination), "now"
    )
    assert isinstance(missing, Err)
    assert isinstance(
        export_combined_report(workshop.coordinator, (), None, str(destination), "now"), Err
    )
