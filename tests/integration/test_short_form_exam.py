"""A scanned 40-question form, end to end: every book stops at Q40 and imports again."""

from __future__ import annotations

from hashlib import sha256
from pathlib import Path
from uuid import uuid4

import openpyxl
from openpyxl import load_workbook

from omr_grader.application.answer_key_use_case import AnswerKeyWorkbookUseCase
from omr_grader.application.dto import RegradeCommand
from omr_grader.application.grading_use_case import GradingUseCase
from omr_grader.domain.enums import AnswerStatus
from omr_grader.domain.errors import Err, Ok
from omr_grader.infrastructure.grading_runtime import CommittedGradingSnapshotReader
from omr_grader.workbooks.answer_key import ANSWER_KEY_HEADERS
from omr_grader.workbooks.response_import import parse_response_book
from omr_grader.workbooks.schemas import RESPONSE_SHEET_NAME, response_headers
from tests.helpers.graded_exam import graded_exam


def _grading(exam) -> GradingUseCase:
    return GradingUseCase(
        CommittedGradingSnapshotReader(exam.coordinator), AnswerKeyWorkbookUseCase(), exam.coordinator
    )


def test_a_forty_question_exam_gets_forty_question_books(tmp_path: Path) -> None:
    exam = graded_exam(tmp_path, "forty", questions=40)

    (score_path,) = exam.folder.glob("*_채점결과_*.xlsx")
    book = load_workbook(score_path)
    for title in ("채점결과", "결과OX"):
        assert book[title].max_column == 46
        assert book[title].cell(1, 45).value == "Q40"
        assert book[title].cell(1, 46).value == "비고"
    assert book["채점결과"].cell(2, 4).value == 40
    assert tuple(cell.value for cell in book[RESPONSE_SHEET_NAME][1]) == response_headers(40)
    assert book["판독 근거"].max_column == 44
    assert book["정답표"].max_row == 41
    (response_path,) = (
        path for path in exam.folder.rglob("*_응답결과_*.xlsx") if "generations" in path.parts
    )
    responses = load_workbook(response_path)[RESPONSE_SHEET_NAME]
    assert tuple(cell.value for cell in responses[1]) == response_headers(40)

    with score_path.open("rb") as source:
        parsed = parse_response_book(
            source,
            sheet_name=RESPONSE_SHEET_NAME,
            session_id="again",
            source_sha256=sha256(score_path.read_bytes()).hexdigest(),
        )
    assert isinstance(parsed, Ok)
    (row,) = parsed.value
    assert all(answer.status is AnswerStatus.NORMAL for answer in row.answers[:40])
    assert all(answer.status is AnswerStatus.UNASKED for answer in row.answers[40:])


def test_a_key_longer_than_the_form_is_refused_with_the_last_printed_question(
    tmp_path: Path,
) -> None:
    exam = graded_exam(tmp_path, "forty", questions=40)
    grading = _grading(exam)
    assert grading.form_questions(exam.session_id, exam.revision) == Ok(tuple(range(1, 41)))
    longer = tmp_path / "inputs" / "key50.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "정답표"
    sheet.append(list(ANSWER_KEY_HEADERS))
    for question in range(1, 51):
        sheet.append([question, "1", 1])
    workbook.save(longer)

    refused = grading.regrade(
        RegradeCommand(exam.session_id, exam.revision, str(longer), "정답표", uuid4().hex)
    )

    assert isinstance(refused, Err)
    assert refused.errors[0].code == "ANSWER_KEY_NOT_ON_FORM"
    assert str(refused.errors[0].context["reason"]).startswith(
        "정답표에 41~50번이 있지만 답안지에는 40번까지만 있습니다."
    )
