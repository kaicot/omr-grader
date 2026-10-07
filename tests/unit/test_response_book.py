from __future__ import annotations

from collections.abc import Iterable
from pathlib import Path

import pytest
from openpyxl import load_workbook

from omr_grader.domain.enums import AnswerStatus, SourceKind, StudentIdStatus
from omr_grader.domain.models import AnswerValue, EffectiveResponse, ImportedResponseRef
from omr_grader.workbooks.response_book import (
    join_notes,
    recognition_note,
    review_note,
    write_effective_response_projection,
    write_response_projection,
)
from omr_grader.workbooks.schemas import RESPONSE_HEADERS, RESPONSE_SHEET_NAME

# Fill and font colors (ARGB) of a cell that still needs a person's check.
YELLOW = ("FFFFEB9C", "FF9C5700")


def _painted(sheet, row: int) -> dict[int, tuple[str, str]]:
    """Fill and font colors of every painted cell of a row, by column number."""
    return {
        column: (cell.fill.fgColor.rgb, cell.font.color.rgb)
        for column in range(1, sheet.max_column + 1)
        if (cell := sheet.cell(row, column)).fill.fill_type is not None
    }


def _answers(changes: dict[int, AnswerValue]) -> tuple[AnswerValue, ...]:
    """Q1..Q100 left blank except for the numbered changes."""
    blank = AnswerValue((), AnswerStatus.BLANK)
    return tuple(changes.get(number, blank) for number in range(1, 101))


def _uncertain(*choices: int) -> AnswerValue:
    return AnswerValue(tuple(choices), AnswerStatus.UNCERTAIN)


# Q3 and Q17 are unconfirmed with a candidate, Q5 is unconfirmed with nothing read.
_MARKED = {
    1: AnswerValue((1,), AnswerStatus.NORMAL),
    3: _uncertain(2),
    4: AnswerValue((1, 3), AnswerStatus.MULTIPLE),
    5: _uncertain(),
    17: _uncertain(4),
    100: AnswerValue((5,), AnswerStatus.NORMAL),
}
# Sheet columns of Q3, Q5 and Q17: Q1 sits in column 5.
_MARKED_COLUMNS = {7: YELLOW, 9: YELLOW, 21: YELLOW}


def _imported(
    serial: int,
    answers: tuple[AnswerValue, ...],
    note: str = "",
    student_id: str = "20240001",
) -> ImportedResponseRef:
    return ImportedResponseRef(
        1,
        f"work-{serial}",
        "a" * 64,
        "recognition",
        serial + 1,
        serial - 1,
        serial,
        f"scan-{serial}.png",
        student_id,
        "미등록",
        answers,
        note,
    )


def _effective(
    serial: int,
    answers: tuple[AnswerValue, ...],
    corrected: tuple[str, ...] = (),
    student_id: str = "20240001",
) -> EffectiveResponse:
    return EffectiveResponse(
        f"work-{serial}",
        SourceKind.IMAGE,
        f"scan-{serial}.png",
        student_id,
        StudentIdStatus.NORMAL,
        answers,
        corrected,
    )


@pytest.mark.parametrize(
    ("questions", "expected"),
    (
        ((), ""),
        ((3,), "확인 필요: 3번"),
        ((17, 3), "확인 필요: 3, 17번"),
        ((3, 3, 17), "확인 필요: 3, 17번"),
        ((3, 4), "확인 필요: 3, 4번"),
        ((3, 4, 5), "확인 필요: 3~5번"),
        ((3, 4, 5, 17), "확인 필요: 3~5, 17번"),
        ((1, 2, 3, 5, 6, 8, 9, 10), "확인 필요: 1~3, 5, 6, 8~10번"),
        (range(1, 101), "확인 필요: 1~100번"),
    ),
)
def test_review_note_names_questions_and_ranges_runs_of_three(
    questions: Iterable[int], expected: str
) -> None:
    assert review_note(questions) == expected


def test_join_notes_keeps_only_the_notes_that_exist() -> None:
    assert join_notes() == ""
    assert join_notes("", "") == ""
    assert join_notes("중복확인필요", "") == "중복확인필요"
    assert join_notes("", "확인 필요: 3번") == "확인 필요: 3번"
    assert join_notes("중복확인필요", "확인 필요: 3번") == "중복확인필요 / 확인 필요: 3번"


@pytest.mark.parametrize(
    ("answers", "student_id_valid", "expected"),
    (
        (_answers({1: AnswerValue((1,), AnswerStatus.NORMAL)}), True, ""),
        (_answers(_MARKED), True, "확인 필요: 3, 5, 17번"),
        (_answers({3: _uncertain(2), 17: _uncertain(4)}), True, "확인 필요: 3, 17번"),
        (
            _answers({3: _uncertain(2), 4: _uncertain(1), 5: _uncertain(), 17: _uncertain(4)}),
            True,
            "확인 필요: 3~5, 17번",
        ),
        (_answers({4: AnswerValue((1, 3), AnswerStatus.MULTIPLE)}), True, ""),
        (_answers({}), False, "학번 확인 필요"),
        (
            _answers({3: _uncertain(2), 17: _uncertain(4)}),
            False,
            "확인 필요: 3, 17번 / 학번 확인 필요",
        ),
        (
            tuple(_uncertain(1) for _ in range(100)),
            False,
            "확인 필요: 1~100번 / 학번 확인 필요",
        ),
    ),
)
def test_recognition_note_reports_unconfirmed_answers_then_unusable_student_id(
    answers: tuple[AnswerValue, ...], student_id_valid: bool, expected: str
) -> None:
    assert recognition_note(answers, student_id_valid=student_id_valid) == expected


def _expected_text(answers: tuple[AnswerValue, ...]) -> list[str | None]:
    """What a reader sees in Q1..Q100: the chosen digits, nothing for no choice."""
    return ["".join(str(choice) for choice in answer.choices) or None for answer in answers]


def _assert_layout(sheet, *, rows: int) -> None:
    """Header, size, frozen identity columns and bold header of a response sheet."""
    assert tuple(cell.value for cell in sheet[1]) == RESPONSE_HEADERS
    assert (sheet.max_row, sheet.max_column) == (rows + 1, 105)
    assert sheet.freeze_panes == "E2"
    assert all(cell.font.b for cell in sheet[1])
    assert not any(cell.font.b for row in sheet.iter_rows(min_row=2) for cell in row)


def test_response_projection_paints_unconfirmed_answers_and_keeps_every_value(
    tmp_path: Path,
) -> None:
    first = _imported(1, _answers(_MARKED), "확인 필요: 3, 5, 17번")
    second = _imported(2, _answers({1: AnswerValue((2,), AnswerStatus.NORMAL)}), student_id="")
    path = tmp_path / "responses.xlsx"

    # Input order differs from the output order, which follows input_ordinal.
    write_response_projection(
        path, (second, first), session_id="session", revision=1, manifest_sha256="0" * 64
    )

    book = load_workbook(path)
    assert book.sheetnames == [RESPONSE_SHEET_NAME]
    sheet = book[RESPONSE_SHEET_NAME]
    _assert_layout(sheet, rows=2)
    for row, source in ((2, first), (3, second)):
        assert sheet.cell(row, 1).value == source.serial
        assert sheet.cell(row, 2).value == source.source_filename
        assert sheet.cell(row, 3).value == (source.raw_student_id or None)
        assert sheet.cell(row, 4).value == source.name
        assert [sheet.cell(row, column).value for column in range(5, 105)] == _expected_text(
            source.answers
        )
        assert sheet.cell(row, 105).value == (source.note or None)
    assert _painted(sheet, 2) == _MARKED_COLUMNS
    assert _painted(sheet, 3) == {}


def test_effective_projection_paints_unconfirmed_answers_and_keeps_manual_note(
    tmp_path: Path,
) -> None:
    corrected = _effective(1, _answers(_MARKED), ("answer_cell:1",))
    untouched = _effective(2, _answers({1: AnswerValue((2,), AnswerStatus.NORMAL)}))
    path = tmp_path / "effective.xlsx"

    write_effective_response_projection(
        path,
        (corrected, untouched),
        session_id="session",
        revision=2,
        manifest_sha256="0" * 64,
        names_by_student_id={"20240001": "가명"},
    )

    sheet = load_workbook(path)[RESPONSE_SHEET_NAME]
    _assert_layout(sheet, rows=2)
    assert [sheet.cell(2, column).value for column in range(5, 105)] == _expected_text(
        corrected.answers
    )
    assert (sheet.cell(2, 4).value, sheet.cell(2, 105).value) == ("가명", "수동 수정 반영")
    assert sheet.cell(3, 105).value is None
    assert _painted(sheet, 2) == _MARKED_COLUMNS
    assert _painted(sheet, 3) == {}
