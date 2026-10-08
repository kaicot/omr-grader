"""The subject-configuration workbook: strict reader and sample."""

from __future__ import annotations

import io
from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest

from omr_grader.domain.errors import Err, Ok
from omr_grader.workbooks.subject_config import (
    SubjectRange,
    parse_subject_config,
    parse_subject_config_bytes,
    subject_config_sample_bytes,
)

_HEADER = ["과목명", "파트", "시작문항", "끝문항"]


def _write(
    tmp_path: Path,
    subject_rows: list[list[object]],
    criteria_rows: list[list[object]] | None = None,
    header: list[str] | None = None,
) -> str:
    book = openpyxl.Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = "과목구성"
    sheet.append(header or _HEADER)
    for row in subject_rows:
        sheet.append(row)
    if criteria_rows is not None:
        criteria = book.create_sheet("합격기준")
        criteria.append(["항목", "기준(%)"])
        for row in criteria_rows:
            criteria.append(row)
    path = tmp_path / "config.xlsx"
    book.save(path)
    return str(path)


def _reasons(result: object) -> str:
    assert isinstance(result, Err)
    return " | ".join(str(error.context["reason"]) for error in result.errors)


def test_a_valid_config_merges_rows_with_the_same_subject(tmp_path):
    path = _write(
        tmp_path,
        [["법규", "파트1", 91, 100], ["법규", 2, 1, 10], ["의학", 1, 1, 90]],
        [["총점", 60], ["과목별", 40.5]],
    )

    result = parse_subject_config(path)

    assert isinstance(result, Ok)
    law, medicine = result.value.subjects
    assert (law.name, medicine.name) == ("법규", "의학")
    assert law.ranges == (SubjectRange(1, 91, 100), SubjectRange(2, 1, 10))
    assert result.value.criteria.total_percent == Decimal(60)
    assert result.value.criteria.per_subject_percent == Decimal("40.5")
    assert result.value.part_numbers() == (1, 2)


def test_criteria_sheet_and_each_criterion_are_optional(tmp_path):
    none = parse_subject_config(_write(tmp_path, [["법규", 1, 1, 10]]))
    one = parse_subject_config(_write(tmp_path, [["법규", 1, 1, 10]], [["과목별", 50]]))

    assert isinstance(none, Ok) and not none.value.criteria.present
    assert isinstance(one, Ok)
    assert one.value.criteria.total_percent is None
    assert one.value.criteria.per_subject_percent == Decimal(50)


def test_extra_columns_and_column_order_are_ignored(tmp_path):
    path = _write(
        tmp_path,
        [["메모", 3, "법규", 1, 5, 6]],
        header=["x", "파트", "과목명", "시작문항", "끝문항", "비고"],
    )

    result = parse_subject_config(path)

    assert isinstance(result, Ok)
    assert result.value.subjects[0].ranges == (SubjectRange(3, 1, 5),)


@pytest.mark.parametrize(
    ("row", "fragment"),
    [
        (["법규", "가", 1, 5], "파트"),
        (["법규", 0, 1, 5], "파트"),
        (["법규", 11, 1, 5], "파트"),
        (["법규", 1, 9, 5], "시작문항이 끝문항보다"),
        (["법규", 1, 0, 5], "1~100"),
        (["법규", 1, 1, 101], "1~100"),
        (["법규", 1, "a", 5], "숫자"),
        (["법규", 1, 1.5, 5], "숫자"),
        ([None, 1, 1, 5], "과목명"),
        (["총점", 1, 1, 5], "쓸 수 없는"),
    ],
)
def test_each_bad_row_is_refused_with_a_korean_reason(tmp_path, row, fragment):
    assert fragment in _reasons(parse_subject_config(_write(tmp_path, [row])))


def test_overlapping_ranges_in_one_part_are_refused_even_across_subjects(tmp_path):
    result = parse_subject_config(
        _write(tmp_path, [["법규", 1, 1, 10], ["의학", 1, 10, 20], ["약학", 2, 1, 20]])
    )

    assert "겹칩니다" in _reasons(result)
    assert isinstance(result, Err) and len(result.errors) == 1


def test_all_problems_are_reported_together(tmp_path):
    result = parse_subject_config(_write(tmp_path, [["법규", 1, 9, 5], ["의학", 0, 1, 5]]))

    assert isinstance(result, Err) and len(result.errors) == 2


@pytest.mark.parametrize(
    ("criteria", "fragment"),
    [
        ([["합계", 60]], "총점"),
        ([["총점", 60], ["총점", 70]], "두 번"),
        ([["총점", "60"]], "숫자"),
        ([["총점", 101]], "0~100"),
        ([["총점", -1]], "0~100"),
        ([["총점", None]], "숫자"),
    ],
)
def test_bad_pass_criteria_are_refused(tmp_path, criteria, fragment):
    path = _write(tmp_path, [["법규", 1, 1, 10]], criteria)

    assert fragment in _reasons(parse_subject_config(path))


def test_missing_sheet_missing_header_and_empty_config_are_refused(tmp_path):
    book = openpyxl.Workbook()
    path = tmp_path / "other.xlsx"
    book.save(path)

    assert "시트가 없습니다" in _reasons(parse_subject_config(str(path)))
    assert "'끝문항' 열" in _reasons(
        parse_subject_config(_write(tmp_path, [], header=["과목명", "파트", "시작문항"]))
    )
    assert "과목이 없습니다" in _reasons(parse_subject_config(_write(tmp_path, [])))


def test_formulas_and_garbage_are_refused(tmp_path):
    path = _write(tmp_path, [["법규", 1, 1, "=1+1"]])

    assert isinstance(parse_subject_config(path), Err)
    assert isinstance(parse_subject_config_bytes(b"not a workbook"), Err)
    assert isinstance(parse_subject_config(str(tmp_path / "missing.xlsx")), Err)


def test_the_sample_is_a_two_part_graduation_exam_of_seven_subjects():
    result = parse_subject_config_bytes(subject_config_sample_bytes())

    assert isinstance(result, Ok)
    subjects = {subject.name: subject.ranges for subject in result.value.subjects}
    assert len(subjects) == 7
    # The law subject spans the end of part 1 and the start of part 2.
    assert subjects["보건의료관계법규"] == (SubjectRange(1, 91, 100), SubjectRange(2, 1, 10))
    assert subjects["해부생리학"] == (SubjectRange(1, 1, 30),)
    assert result.value.part_numbers() == (1, 2)
    assert result.value.criteria.total_percent == Decimal(60)
    assert result.value.criteria.per_subject_percent == Decimal(60)
    book = openpyxl.load_workbook(io.BytesIO(subject_config_sample_bytes()))
    assert book.sheetnames == ["과목구성", "합격기준", "설명"]
