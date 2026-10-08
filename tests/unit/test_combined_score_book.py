"""The 합산 성적표 builder: totals, ranks, review rows, subjects, and pass/fail."""

from __future__ import annotations

import io
from decimal import Decimal

import openpyxl
import pytest

from omr_grader.workbooks.combined_score_book import (
    PartScores,
    PartStudent,
    build_combined_score_book,
    summarize_combined_scores,
)
from omr_grader.workbooks.subject_config import (
    PassCriteria,
    Subject,
    SubjectConfig,
    SubjectRange,
)

_YELLOW = "FFFFEB9C"
_ORANGE = "FFF8CBAD"
_RED_FILL, _RED_FONT = "FFFFC7CE", "FF9C0006"
_GREEN_FILL = "FFC6EFCE"


def _points(asked: int) -> tuple[Decimal | None, ...]:
    """The first ``asked`` questions are worth one point; the rest are not asked."""
    return tuple(Decimal(1) if number <= asked else None for number in range(1, 101))


def _student(
    student_id: str | None,
    name: str,
    right: set[int] | None,
    order: int = 1,
    asked: int = 10,
) -> PartStudent:
    """``right`` are the questions answered correctly; ``None`` holds the sheet back."""
    if right is None:
        return PartStudent(student_id, name, None, (), f"scan.pdf p{order}", order)
    earned = tuple(
        None if number > asked else Decimal(1 if number in right else 0) for number in range(1, 101)
    )
    return PartStudent(student_id, name, Decimal(len(right)), earned, f"scan.pdf p{order}", order)


def _part(label: str, rows: tuple[PartStudent, ...], asked: int = 10) -> PartScores:
    return PartScores(
        label,
        f"{label} 시험",
        f"260101_090000_{label}",
        "2026-01-01T09:30:00Z",
        _points(asked),
        rows,
    )


def _book(parts, subjects=None) -> openpyxl.Workbook:
    data = build_combined_score_book(parts, subjects, "2026-02-01 10:00")
    return openpyxl.load_workbook(io.BytesIO(data))


def _table(book, name="합산결과") -> list[dict[str, object]]:
    rows = list(book[name].iter_rows(values_only=True))
    return [dict(zip(rows[0], row, strict=True)) for row in rows[1:]]


def _two_parts() -> tuple[PartScores, PartScores]:
    first = _part(
        "파트1",
        (
            _student("20260001", "다람", set(range(1, 9)), 1),
            _student("20260002", "가온", set(range(1, 11)), 2),
            _student("20260003", "나래", set(range(1, 9)), 3),
            _student("20260004", "", set(range(1, 6)), 4),
        ),
    )
    second = _part(
        "파트2",
        (
            _student("20260002", "가온", set(range(1, 11)), 1),
            _student("20260001", "다람", set(range(1, 3)), 2),
            _student("20260003", "나래", set(range(1, 3)), 3),
            _student("20260004", "", set(range(1, 6)), 4),
        ),
    )
    return first, second


def test_totals_ranks_and_order_follow_the_per_exam_book():
    book = _book(_two_parts())

    rows = _table(book)

    assert book.sheetnames == ["합산결과", "파트"]
    assert list(rows[0]) == [
        "순번",
        "학번",
        "이름",
        "파트1 점수",
        "파트2 점수",
        "총점",
        "통합 석차",
        "비고",
    ]
    assert book["합산결과"].freeze_panes == "D2"
    assert book["합산결과"]["A1"].font.bold
    # 가나다 by name first, unnamed students last.
    assert [row["이름"] for row in rows] == ["가온", "나래", "다람", None]
    assert [row["순번"] for row in rows] == [1, 2, 3, 4]
    assert [row["총점"] for row in rows] == [20, 10, 10, 10]
    assert [row["통합 석차"] for row in rows] == [1, 2, 2, 2]
    assert rows[0]["파트1 점수"] == 10 and rows[0]["파트2 점수"] == 10
    assert all(row["비고"] in (None, "") for row in rows)


def test_ties_leave_a_gap_after_them():
    part = _part(
        "파트1",
        tuple(
            _student(f"2026000{index}", name, set(range(1, total + 1)), index)
            for index, (name, total) in enumerate((("가", 9), ("나", 7), ("다", 7), ("라", 5)), 1)
        ),
    )

    rows = _table(_book([part]))

    assert [row["통합 석차"] for row in rows] == [1, 2, 2, 4]


def test_a_student_missing_from_a_part_is_yellow_with_a_note():
    first, second = _two_parts()
    second = _part("파트2", second.rows[:2])

    book = _book((first, second))
    rows = _table(book)

    missing = next(row for row in rows if row["이름"] == "나래")
    assert missing["총점"] is None and missing["통합 석차"] is None
    assert missing["파트1 점수"] == 8 and missing["파트2 점수"] is None
    assert missing["비고"] == "확인 필요: 파트2 기록 없음 (학번 확인)"
    sheet = book["합산결과"]
    assert all(cell.fill.start_color.rgb == _YELLOW for cell in sheet[3])
    assert sheet[3][0].font.color.rgb == "FF9C5700"
    assert sheet[2][0].fill.start_color.rgb != _YELLOW
    # Ranks count only complete rows.
    assert [row["통합 석차"] for row in rows if row["총점"] is not None] == [1, 2]
    assert summarize_combined_scores((first, second)).needs_review == 2


def test_a_held_back_total_blocks_the_combined_total():
    first, second = _two_parts()
    first = _part("파트1", (_student("20260002", "가온", None, 1), *first.rows[:1]))

    rows = _table(_book((first, second)))

    held = next(row for row in rows if row["이름"] == "가온")
    assert held["총점"] is None and held["통합 석차"] is None
    assert held["파트1 점수"] is None and held["파트2 점수"] == 10
    assert held["비고"] == "확인 필요: 파트1 확인 필요 문항"


def test_every_problem_is_listed_in_the_note():
    first = _part("파트1", (_student("20260002", "가온", None, 1),))
    second = _part("파트2", ())

    rows = _table(_book((first, second)))

    assert rows[0]["비고"] == "확인 필요: 파트1 확인 필요 문항, 파트2 기록 없음 (학번 확인)"


def test_duplicate_ids_inside_a_part_need_review():
    first = _part(
        "파트1",
        (
            _student("20260002", "가온", {1, 2}, 1),
            _student("20260002", "가온", {1, 2, 3}, 2),
            _student("20260001", "나래", {1}, 3),
        ),
    )
    second = _part(
        "파트2", (_student("20260002", "가온", {1}, 1), _student("20260001", "나래", {1}, 2))
    )

    rows = _table(_book((first, second)))

    assert len(rows) == 2
    duplicate = rows[0]
    assert duplicate["이름"] == "가온" and duplicate["총점"] is None
    assert duplicate["비고"] == "확인 필요: 파트1 학번 중복"
    assert rows[1]["총점"] == 2 and rows[1]["통합 석차"] == 1


def test_unreadable_ids_are_listed_on_their_own_sheet():
    first, second = _two_parts()
    first = _part("파트1", (*first.rows, _student(None, "", {1}, 5), _student("  ", "", None, 6)))
    second = _part("파트2", (*second.rows, _student(None, "", {1, 2}, 5)))

    book = _book((first, second))

    assert book.sheetnames == ["합산결과", "학번 확인 필요", "파트"]
    rows = _table(book, "학번 확인 필요")
    assert [(row["파트"], row["순번"], row["원본"], row["점수"]) for row in rows] == [
        ("파트1", 5, "scan.pdf p5", 1),
        ("파트1", 6, "scan.pdf p6", None),
        ("파트2", 5, "scan.pdf p5", 2),
    ]
    assert len(_table(book)) == 4
    counts = summarize_combined_scores((first, second))
    assert (counts.students, counts.complete, counts.unreadable) == (4, 4, 3)


def test_no_unreadable_sheet_and_the_parts_sheet_describes_each_session():
    book = _book(_two_parts())

    assert "학번 확인 필요" not in book.sheetnames
    rows = _table(book, "파트")
    assert rows[0]["파트"] == "파트1" and rows[0]["폴더"] == "260101_090000_파트1"
    assert rows[0]["학생 수"] == 4 and rows[0]["만점"] == 10
    assert rows[0]["채점 시각"] == "2026-01-01T09:30:00Z"
    assert book["파트"].cell(book["파트"].max_row, 2).value == "2026-02-01 10:00"


def _config(**criteria) -> SubjectConfig:
    return SubjectConfig(
        (
            Subject("법규", (SubjectRange(1, 6, 10), SubjectRange(2, 1, 5))),
            Subject("의학", (SubjectRange(1, 1, 5), SubjectRange(2, 6, 10))),
        ),
        PassCriteria(**criteria),
    )


def test_subject_points_and_maxima_merge_ranges_across_parts():
    first = _part("파트1", (_student("20260001", "가", {1, 2, 6, 7, 8}, 1),))
    second = _part("파트2", (_student("20260001", "가", {1, 6, 7, 8, 9}, 1),))

    book = _book((first, second), _config())

    row = _table(book)[0]
    # 법규 = part 1 Q6-10 (3 right) + part 2 Q1-5 (1 right); 의학 = part 1 Q1-5 (2) + part 2 Q6-10 (4).
    assert (row["법규"], row["의학"], row["총점"]) == (4, 6, 10)
    assert "합격 여부" not in row
    assert _table(book, "과목구성") == [
        {"과목명": "법규", "범위": "파트1 6~10번, 파트2 1~5번", "만점": 10},
        {"과목명": "의학", "범위": "파트1 1~5번, 파트2 6~10번", "만점": 10},
    ]


def test_questions_the_key_does_not_ask_add_no_maximum():
    first = _part("파트1", (_student("20260001", "가", {1, 2, 3}, 1, asked=4),), asked=4)
    second = _part("파트2", (_student("20260001", "가", {1}, 1, asked=4),), asked=4)

    book = _book((first, second), _config())

    # 법규 asks part 1 Q6-10 (none) and part 2 Q1-5 (4); 의학 asks part 1 Q1-5 (4) and none.
    assert [row["만점"] for row in _table(book, "과목구성")] == [4, 4]


_HALF = {1, 2, 3, 4, 5}


@pytest.mark.parametrize(
    ("criteria", "right", "expected"),
    [
        # Total maximum is 20: 60% needs 12.
        ({"total_percent": Decimal(60)}, ({1, 2, 3, 4, 5, 6}, {1, 2, 3, 4, 5, 6}), "합격"),
        ({"total_percent": Decimal(60)}, ({1, 2, 3, 4, 5, 6}, {1, 2, 3, 4, 5}), "불합격"),
        # Each subject has a maximum of 10: 40% needs 4 in both.
        # 법규 = part 1 Q6-10 + part 2 Q1-5, 의학 = part 1 Q1-5 + part 2 Q6-10.
        ({"per_subject_percent": Decimal(40)}, ({1, 2, 6, 7}, {1, 2, 6, 7}), "합격"),
        ({"per_subject_percent": Decimal(60)}, ({1, 2, 6, 7}, {1, 2, 6, 7}), "불합격"),
        # Half the total, but nothing in 법규.
        ({"per_subject_percent": Decimal(40)}, (_HALF, {6, 7, 8, 9, 10}), "불합격"),
        # Both criteria: a failing subject sinks a high total, a low total sinks good subjects.
        (
            {"total_percent": Decimal(50), "per_subject_percent": Decimal(40)},
            (_HALF, {6, 7, 8, 9, 10}),
            "불합격",
        ),
        (
            {"total_percent": Decimal(80), "per_subject_percent": Decimal(40)},
            ({1, 2, 6, 7}, {1, 2, 6, 7}),
            "불합격",
        ),
        (
            {"total_percent": Decimal(50), "per_subject_percent": Decimal(40)},
            ({1, 2, 3, 6, 7, 8}, {1, 2, 6, 7}),
            "합격",
        ),
    ],
)
def test_pass_fail_applies_each_criterion(criteria, right, expected):
    first = _part("파트1", (_student("20260001", "가", right[0], 1),))
    second = _part("파트2", (_student("20260001", "가", right[1], 1),))

    book = _book((first, second), _config(**criteria))

    row = _table(book)[0]
    assert row["합격 여부"] == expected
    headers = [cell.value for cell in book["합산결과"][1]]
    assert headers[-2:] == ["합격 여부", "비고"]
    assert book["과목구성"].cell(5, 1).value == "합격기준"


def test_an_incomplete_row_has_no_subjects_or_verdict():
    first, second = _two_parts()
    second = _part("파트2", second.rows[:2])

    rows = _table(_book((first, second), _config(total_percent=Decimal(10))))

    missing = next(row for row in rows if row["이름"] == "나래")
    assert missing["법규"] is None and missing["의학"] is None
    assert missing["합격 여부"] is None
    assert str(missing["비고"]).startswith("확인 필요:")


def test_boundary_percentages_are_exact():
    part = _part("파트1", (_student("20260001", "가", set(range(1, 8)), 1),))
    config = SubjectConfig(
        (Subject("가", (SubjectRange(1, 1, 10),)),), PassCriteria(total_percent=Decimal(70))
    )

    # 7 of 10 is exactly 70%.
    assert _table(_book([part], config))[0]["합격 여부"] == "합격"


def test_config_for_a_part_that_is_not_selected_is_refused():
    first, _ = _two_parts()

    with pytest.raises(ValueError, match="part"):
        build_combined_score_book([first], _config(), "now")
    with pytest.raises(ValueError):
        build_combined_score_book([], None, "now")


def _subject_rows(book, name="과목별 합격", first=4) -> list[dict[str, object]]:
    sheet = book[name]
    headers = [cell.value for cell in sheet[1]]
    return [
        dict(zip(headers, (cell.value for cell in sheet[row]), strict=True))
        for row in range(first, sheet.max_row + 1)
        if isinstance(sheet.cell(row, 1).value, int)
    ]


def _graduation() -> tuple[tuple[PartScores, PartScores], SubjectConfig]:
    """Three students over two ten-question parts; 법규 spans both parts."""
    first = _part(
        "파트1",
        (
            _student("20260001", "가", set(range(1, 11)), 1),
            _student("20260002", "나", {1, 2, 6, 7, 8, 9, 10}, 2),
            _student("20260003", "다", {1, 2, 3, 6, 7}, 3),
        ),
    )
    second = _part(
        "파트2",
        (
            _student("20260001", "가", set(range(1, 11)), 1),
            _student("20260002", "나", set(range(1, 11)), 2),
            _student("20260003", "다", {1, 2, 3, 6, 7}, 3),
        ),
    )
    config = SubjectConfig(
        (
            Subject("해부", (SubjectRange(1, 1, 5),)),
            Subject("법규", (SubjectRange(1, 6, 10), SubjectRange(2, 1, 5))),
            Subject("평가", (SubjectRange(2, 6, 10),)),
        ),
        PassCriteria(total_percent=Decimal(60), per_subject_percent=Decimal(60)),
    )
    return (first, second), config


def test_the_subject_sheet_comes_first_with_minimums_and_colored_verdicts():
    parts, config = _graduation()

    book = _book(parts, config)

    assert book.sheetnames == ["과목별 합격", "합산결과", "파트", "과목구성"]
    assert book.active.title == "과목별 합격"
    assert [sheet.sheet_view.tabSelected for sheet in book.worksheets] == [
        True,
        False,
        False,
        False,
    ]
    sheet = book["과목별 합격"]
    assert [cell.value for cell in sheet[1]] == [
        "순번", "학번", "이름", "해부", "법규", "평가", "합계",
        "미달 과목 수", "판정", "미달 과목", "비고",
    ]  # fmt: skip
    assert [sheet.cell(2, column).value for column in range(1, 8)] == [
        "만점", None, None, 5, 10, 5, 20,
    ]  # fmt: skip
    # 60% of 5 is 3, of 10 is 6, of 20 is 12.
    assert [sheet.cell(3, column).value for column in range(1, 8)] == [
        "합격 최소 점수", None, None, 3, 6, 3, 12,
    ]  # fmt: skip
    assert sheet.freeze_panes == "D4"
    rows = _subject_rows(book)
    assert [(row["이름"], row["해부"], row["법규"], row["평가"], row["합계"]) for row in rows] == [
        ("가", 5, 10, 5, 20),
        ("나", 2, 10, 5, 17),
        ("다", 3, 5, 2, 10),
    ]
    assert [(row["미달 과목 수"], row["판정"], row["미달 과목"]) for row in rows] == [
        (0, "합격", None),
        (1, "불합격", "해부"),
        (2, "불합격", "법규, 평가, 총점"),
    ]
    passed, failed = sheet["I4"], sheet["I5"]
    assert passed.fill.start_color.rgb == _GREEN_FILL
    assert failed.fill.start_color.rgb == _RED_FILL and failed.font.color.rgb == _RED_FONT
    assert sheet["J5"].font.color.rgb == _RED_FONT
    # Only below-minimum scores are orange; the short total is orange too.
    assert sheet["D5"].fill.start_color.rgb == _ORANGE
    assert sheet["E5"].fill.start_color.rgb != _ORANGE
    assert [sheet.cell(6, column).fill.start_color.rgb == _ORANGE for column in range(4, 8)] == [
        False,
        True,
        True,
        True,
    ]
    # The verdict on 합산결과 carries the same colors.
    verdicts = book["합산결과"]
    column = [cell.value for cell in verdicts[1]].index("합격 여부") + 1
    assert verdicts.cell(2, column).fill.start_color.rgb == _GREEN_FILL
    assert verdicts.cell(3, column).font.color.rgb == _RED_FONT


def test_the_subject_sheet_footer_counts_passes_and_averages():
    parts, config = _graduation()

    sheet = _book(parts, config)["과목별 합격"]

    labels = {sheet.cell(row, 1).value: row for row in range(1, sheet.max_row + 1)}
    passes, means, rates = (
        labels["과목 통과 인원"],
        labels["과목 평균"],
        labels["과목 평균 정답률"],
    )
    assert passes == 8  # one blank row after the three students
    assert [sheet.cell(passes, column).value for column in range(4, 7)] == [2, 2, 2]
    assert (sheet.cell(passes, 8).value, sheet.cell(passes, 9).value) == ("합격 인원", 1)
    assert [sheet.cell(means, column).value for column in range(4, 8)] == pytest.approx(
        [10 / 3, 25 / 3, 4, 47 / 3]
    )
    assert sheet.cell(means, 4).number_format == "0.0"
    assert sheet.cell(rates, 5).value == pytest.approx(25 / 30)
    assert sheet.cell(rates, 5).number_format == "0%"
    notes = [
        str(sheet.cell(row, 1).value)
        for row in range(rates + 1, sheet.max_row + 1)
        if sheet.cell(row, 1).value
    ]
    assert notes[0] == (
        "과목 구성: 해부 파트1 1~5번; 법규 파트1 6~10번 + 파트2 1~5번; 평가 파트2 6~10번"
    )
    assert notes[1].startswith("합격 기준: 모든 과목 60% 이상 · 총점 60% 이상.")


def test_a_student_needing_review_has_no_scores_and_is_left_out_of_the_footer():
    (first, second), config = _graduation()
    second = _part("파트2", second.rows[:2])

    book = _book((first, second), config)

    rows = _subject_rows(book)
    review = rows[2]
    assert review["이름"] == "다" and review["판정"] == "확인 필요"
    assert review["합계"] is None and review["해부"] is None
    assert review["비고"] == "확인 필요: 파트2 기록 없음 (학번 확인)"
    sheet = book["과목별 합격"]
    assert all(cell.fill.start_color.rgb == _YELLOW for cell in sheet[6])
    passes = next(
        row for row in range(1, sheet.max_row + 1) if sheet.cell(row, 1).value == "과목 통과 인원"
    )
    assert [sheet.cell(passes, column).value for column in range(4, 7)] == [1, 2, 2]


def test_without_criteria_the_subject_sheet_only_scores():
    parts, config = _graduation()
    plain = SubjectConfig(config.subjects)

    book = _book(parts, plain)

    assert book.sheetnames[0] == "과목별 점수"
    sheet = book["과목별 점수"]
    assert [cell.value for cell in sheet[1]][-2:] == ["합계", "비고"]
    assert sheet.cell(3, 1).value == 1 and sheet.freeze_panes == "D3"
    assert all(cell.fill.start_color.rgb != _ORANGE for cell in sheet[3])


def test_fractional_points_show_the_exact_minimum():
    points = tuple(Decimal("2.5") if number <= 4 else None for number in range(1, 101))
    student = PartStudent(
        "20260001",
        "가",
        Decimal("7.5"),
        tuple(
            Decimal("2.5") if number <= 3 else Decimal(0) if number == 4 else None
            for number in range(1, 101)
        ),
        "scan.pdf p1",
        1,
    )
    part = PartScores("파트1", "시험", "folder", None, points, (student,))
    config = SubjectConfig(
        (Subject("가", (SubjectRange(1, 1, 4),)),), PassCriteria(per_subject_percent=Decimal(55))
    )

    sheet = _book([part], config)["과목별 합격"]

    # 55% of 10 is 5.5; scores move in steps of 2.5, so the minimum is shown as is.
    assert sheet["D3"].value == 5.5
    assert sheet["G4"].value == "합격"


def test_text_that_looks_like_a_formula_is_stored_as_text():
    part = _part("파트1", (_student("20260001", "=SUM(A1)", {1}, 1),))

    data = build_combined_score_book([part], None, "now")

    book = openpyxl.load_workbook(io.BytesIO(data))
    assert book["합산결과"]["C2"].data_type == "s"
    assert not str(book["합산결과"]["C2"].value).startswith("=")
