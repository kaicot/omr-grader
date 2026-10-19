"""Workbooks whose sheets do not record their used range still import.

Excel always writes a ``<dimension>`` element; some other tools leave it out, and openpyxl's
read-only sheets then report no size. A roster made that way stopped every scan.
"""

from __future__ import annotations

import io
import re
import zipfile
from hashlib import sha256
from pathlib import Path
from uuid import uuid4

import pytest
from openpyxl import Workbook, load_workbook

from omr_grader.application.dto import ScanCommand, ScanSource
from omr_grader.domain.enums import ExamTerm
from omr_grader.domain.errors import Err, Ok, error_site
from omr_grader.infrastructure import scan_runtime
from omr_grader.ingestion.roster import ROSTER_HEADERS, import_roster
from omr_grader.workbooks.answer_key import import_answer_key, write_answer_key_sample
from omr_grader.workbooks.response_import import parse_response_book
from omr_grader.workbooks.schemas import RESPONSE_HEADERS, RESPONSE_SHEET_NAME, sheet_extent


def _without_dimension(path: Path) -> Path:
    """Rewrite ``path`` the way some non-Excel tools save it: no ``<dimension>`` element."""
    source = zipfile.ZipFile(io.BytesIO(path.read_bytes()))
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as target:
        for item in source.infolist():
            payload = source.read(item.filename)
            if item.filename.startswith("xl/worksheets/sheet"):
                payload = re.sub(rb"<dimension[^>]*/>", b"", payload)
            target.writestr(item, payload)
    path.write_bytes(output.getvalue())
    return path


def _roster(path: Path) -> Path:
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = "명단"
    sheet.append(ROSTER_HEADERS)
    sheet.append(["1", "20260001", "가나다"])
    sheet.append(["2", "20260002", "라마바"])
    book.save(path)
    return _without_dimension(path)


def test_sheet_extent_counts_the_used_range_when_the_file_does_not_record_it(tmp_path) -> None:
    path = tmp_path / "book.xlsx"
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.append(["a", "b"])
    sheet["E7"] = "x"
    book.create_sheet("빈 시트")
    book.save(path)
    _without_dimension(path)

    loaded = load_workbook(path, read_only=True)
    first, empty = loaded.worksheets
    assert (first.max_row, first.max_column) == (None, None)

    assert sheet_extent(first) == (7, 5)
    assert sheet_extent(empty) == (0, 0)
    loaded.close()


def test_a_roster_without_a_recorded_size_imports(tmp_path) -> None:
    imported = import_roster(str(_roster(tmp_path / "roster.xlsx")), "명단")

    assert isinstance(imported, Ok)
    assert [row.student_id for row in imported.value.rows] == ["20260001", "20260002"]


def test_an_answer_key_without_a_recorded_size_imports(tmp_path) -> None:
    path = tmp_path / "key.xlsx"
    assert isinstance(write_answer_key_sample(str(path)), Ok)
    _without_dimension(path)

    imported = import_answer_key(str(path), "정답표")

    assert isinstance(imported, Ok)
    assert imported.value.entries


def test_a_response_book_without_a_recorded_size_imports(tmp_path) -> None:
    path = tmp_path / "responses.xlsx"
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = RESPONSE_SHEET_NAME
    sheet.append(list(RESPONSE_HEADERS))
    sheet.append([1, "scan-01.jpg", "20260001", "가나다", "2", *([""] * 99), ""])
    book.save(path)
    _without_dimension(path)

    with path.open("rb") as source:
        parsed = parse_response_book(
            source,
            sheet_name=RESPONSE_SHEET_NAME,
            session_id="import-session",
            source_sha256=sha256(path.read_bytes()).hexdigest(),
        )

    assert isinstance(parsed, Ok)
    assert [row.raw_student_id for row in parsed.value] == ["20260001"]


def _command(roster: Path) -> ScanCommand:
    return ScanCommand(
        f"scan-{uuid4().hex}", uuid4().hex, 0, "시험", None, ExamTerm.UNSPECIFIED,
        "form.omrtemplate", str(roster), ScanSource(("scan.pdf",)), 5, False,
    )


def test_an_unreadable_roster_is_reported_in_korean_with_the_failing_site(
    tmp_path, monkeypatch
) -> None:
    def broken(*_: object) -> object:
        raise TypeError("'>' not supported between instances of 'NoneType' and 'int'")

    monkeypatch.setattr(scan_runtime, "import_roster", broken)
    runtime = scan_runtime.ScanRuntime(None, None)  # type: ignore[arg-type]

    failed = runtime._roster(_command(_roster(tmp_path / "roster.xlsx")))

    assert isinstance(failed, Err)
    error = failed.errors[0]
    assert error.code == "ROSTER_TYPEERROR"
    assert "명단 파일을 읽지 못했습니다" in str(error.context["reason"])
    assert "test_dimensionless_workbooks" not in str(error.context.get("site"))


def test_a_roster_with_wrong_headers_says_what_to_check(tmp_path) -> None:
    path = tmp_path / "roster.xlsx"
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.append(["번호", "학번", "성명"])
    book.save(path)
    runtime = scan_runtime.ScanRuntime(None, None)  # type: ignore[arg-type]

    failed = runtime._roster(_command(path))

    assert isinstance(failed, Err)
    assert failed.errors[0].code == "XLSX_HEADERS_INVALID"
    assert "연번 · 학번 · 이름" in str(failed.errors[0].context["reason"])


def test_error_site_names_program_lines_but_never_the_message() -> None:
    with pytest.raises(AttributeError) as raised:
        sheet_extent(object())

    site = error_site(raised.value)

    assert site.startswith("workbooks/schemas.py:")
    assert "sheet_extent" in site
    assert "object" not in site
