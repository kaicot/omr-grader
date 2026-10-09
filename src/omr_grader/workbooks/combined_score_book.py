"""Formula-free 합산 성적표: per-student totals across the parts of a split exam.

Each part is one separately scanned and graded session.  The builder is pure: it takes
already-read per-part scores and returns the workbook bytes, so the caller decides where and
how to write them.
"""

from __future__ import annotations

from collections.abc import Sequence
from dataclasses import dataclass
from datetime import datetime
from decimal import ROUND_CEILING, Decimal
from io import BytesIO
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.styles.styleable import StyleableObject
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from omr_grader.workbooks.response_book import mark_review, style_header_row
from omr_grader.workbooks.schemas import escape_formula_text
from omr_grader.workbooks.subject_config import Subject, SubjectConfig

QUESTION_COUNT = 100
RESULT_SHEET_NAME = "합산결과"
UNREADABLE_SHEET_NAME = "학번 확인 필요"
# Stored times are UTC; the screens and this report show Korean time.
_LOCAL_ZONE = ZoneInfo("Asia/Seoul")
PARTS_SHEET_NAME = "파트"
SUBJECT_SHEET_NAME = "과목구성"
SUBJECT_PASS_SHEET_NAME = "과목별 합격"
SUBJECT_SCORE_SHEET_NAME = "과목별 점수"
PASS_TEXT = "합격"
FAIL_TEXT = "불합격"
REVIEW_TEXT = "확인 필요"
TOTAL_SHORT_TEXT = "총점"
_FREEZE_PANES = "D2"
_NOTE_SEPARATOR = ", "
_HUNDRED = Decimal(100)


def _solid(color: str) -> PatternFill:
    return PatternFill(fill_type="solid", start_color=color, end_color=color)


# The same colors teachers know from Excel's good / bad cell styles; below-minimum subject
# scores are orange so they never look like the yellow "check this" cells.
_BAND_FILL = _solid("FFD9D9D9")
_BELOW_FILL = _solid("FFF8CBAD")
_PASS_FILL = _solid("FFC6EFCE")
_FAIL_FILL = _solid("FFFFC7CE")
_BOLD = Font(name="Calibri", size=11, bold=True)
_PASS_FONT = Font(name="Calibri", size=11, bold=True, color="FF006100")
_FAIL_FONT = Font(name="Calibri", size=11, bold=True, color="FF9C0006")
_NOTE_FONT = Font(name="Calibri", size=10, color="FF555555")
_CENTER = Alignment(horizontal="center", vertical="center")
_HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center", wrap_text=True)


@dataclass(frozen=True, slots=True)
class PartStudent:
    """One graded answer sheet of a part.

    ``student_id`` is ``None`` (or blank) when the ID could not be read.  ``total`` is ``None``
    when the sheet is held back because answers still need a check; ``earned`` then is empty.
    Otherwise ``earned[question - 1]`` is the points the student got for that question
    (``None`` for questions the key does not ask).  ``part_order`` is the sheet's 순번 in the
    part's own score book.
    """

    student_id: str | None
    name: str
    total: Decimal | None
    earned: tuple[Decimal | None, ...]
    source_label: str
    part_order: int

    @property
    def readable_id(self) -> str | None:
        return self.student_id.strip() or None if self.student_id else None


@dataclass(frozen=True, slots=True)
class PartScores:
    """One graded session.  ``question_points[question - 1]`` is the key's points for a
    question, ``None`` where the key does not ask it (everyone-correct questions count as asked).
    """

    label: str
    exam_name: str
    folder_name: str
    graded_at: str | None
    question_points: tuple[Decimal | None, ...]
    rows: tuple[PartStudent, ...]

    @property
    def maximum(self) -> Decimal:
        return sum((item for item in self.question_points if item is not None), Decimal(0))


@dataclass(frozen=True, slots=True)
class CombinedCounts:
    students: int
    complete: int
    needs_review: int
    unreadable: int


@dataclass(frozen=True, slots=True)
class _Row:
    student_id: str
    name: str
    part_totals: tuple[Decimal | None, ...]
    total: Decimal | None
    rank: int | None
    subject_points: tuple[Decimal | None, ...]
    passed: bool | None
    problems: tuple[str, ...]

    @property
    def complete(self) -> bool:
        return not self.problems


def _text(value: str) -> str:
    return escape_formula_text(value)


def _validate(parts: Sequence[PartScores], subjects: SubjectConfig | None) -> None:
    if not parts:
        raise ValueError("at least one part is required")
    if len({part.label for part in parts}) != len(parts):
        raise ValueError("part labels must be unique")
    for part in parts:
        if len(part.question_points) != QUESTION_COUNT or any(
            item.earned and len(item.earned) != QUESTION_COUNT for item in part.rows
        ):
            raise ValueError("question points must cover 100 questions")
    if subjects is not None and any(number > len(parts) for number in subjects.part_numbers()):
        raise ValueError("subject configuration refers to a part that is not selected")


def _subject_max(subject: Subject, parts: Sequence[PartScores]) -> Decimal:
    total = Decimal(0)
    for item in subject.ranges:
        points = parts[item.part - 1].question_points
        total += sum(
            (value for value in points[item.first - 1 : item.last] if value is not None),
            Decimal(0),
        )
    return total


def _whole_points(subject: Subject, parts: Sequence[PartScores]) -> bool:
    """Whether every asked question of the subject is worth a whole number of points."""
    return all(
        value == value.to_integral_value()
        for item in subject.ranges
        for value in parts[item.part - 1].question_points[item.first - 1 : item.last]
        if value is not None
    )


def _minimum(maximum: Decimal, percent: Decimal, whole: bool) -> Decimal:
    """The lowest score that reaches ``percent`` of ``maximum``."""
    exact = percent * maximum / _HUNDRED
    return exact.to_integral_value(rounding=ROUND_CEILING) if whole else exact.normalize()


def _reaches(earned: Decimal, maximum: Decimal, percent: Decimal) -> bool:
    """The one comparison every verdict uses: ``earned`` is at least ``percent`` of ``maximum``."""
    return earned * _HUNDRED >= percent * maximum


def _subject_earned(subject: Subject, students: Sequence[PartStudent | None]) -> Decimal:
    total = Decimal(0)
    for item in subject.ranges:
        student = students[item.part - 1]
        if student is None:
            continue
        total += sum(
            (value for value in student.earned[item.first - 1 : item.last] if value is not None),
            Decimal(0),
        )
    return total


def _collect(
    parts: Sequence[PartScores], subjects: SubjectConfig | None
) -> tuple[list[_Row], list[tuple[PartScores, PartStudent]]]:
    unreadable: list[tuple[PartScores, PartStudent]] = []
    by_id: list[dict[str, list[PartStudent]]] = []
    for part in parts:
        grouped: dict[str, list[PartStudent]] = {}
        for student in part.rows:
            student_id = student.readable_id
            if student_id is None:
                unreadable.append((part, student))
            else:
                grouped.setdefault(student_id, []).append(student)
        by_id.append(grouped)
    ids = {student_id for grouped in by_id for student_id in grouped}
    config = subjects.subjects if subjects is not None else ()
    maximum = sum((part.maximum for part in parts), Decimal(0))
    rows: list[_Row] = []
    for student_id in ids:
        found = [grouped.get(student_id, []) for grouped in by_id]
        names = [item.name for group in found for item in group if item.name]
        problems: list[str] = []
        singles: list[PartStudent | None] = []
        for part, group in zip(parts, found, strict=True):
            if not group:
                problems.append(f"{part.label} 기록 없음 (학번 확인)")
                singles.append(None)
            elif len(group) > 1:
                problems.append(f"{part.label} 학번 중복")
                singles.append(None)
            elif group[0].total is None:
                problems.append(f"{part.label} 확인 필요 문항")
                singles.append(None)
            else:
                singles.append(group[0])
        part_totals = tuple(group[0].total if len(group) == 1 else None for group in found)
        complete = not problems
        total = (
            sum((item.total for item in singles if item and item.total is not None), Decimal(0))
            if complete
            else None
        )
        points = tuple(
            _subject_earned(subject, singles) if complete else None for subject in config
        )
        passed: bool | None = None
        if complete and subjects is not None and subjects.criteria.present:
            criteria = subjects.criteria
            assert total is not None
            passed = criteria.total_percent is None or _reaches(
                total, maximum, criteria.total_percent
            )
            if criteria.per_subject_percent is not None:
                per_subject = criteria.per_subject_percent
                passed = passed and all(
                    earned is not None
                    and _reaches(earned, _subject_max(subject, parts), per_subject)
                    for subject, earned in zip(config, points, strict=True)
                )
        rows.append(
            _Row(
                student_id,
                names[0] if names else "",
                part_totals,
                total,
                None,
                points,
                passed,
                tuple(problems),
            )
        )
    # Teachers read by name in 가나다 order, then students without a name by student ID.
    rows.sort(key=lambda row: (not row.name, row.name, row.student_id))
    totals = [row.total for row in rows if row.total is not None]
    ranked = [
        _Row(
            row.student_id,
            row.name,
            row.part_totals,
            row.total,
            None if row.total is None else 1 + sum(other > row.total for other in totals),
            row.subject_points,
            row.passed,
            row.problems,
        )
        for row in rows
    ]
    unreadable.sort(key=lambda item: (parts.index(item[0]), item[1].part_order))
    return ranked, unreadable


def summarize_combined_scores(parts: Sequence[PartScores]) -> CombinedCounts:
    """How many students the combined report has and how many of them are complete."""
    rows, unreadable = _collect(parts, None)
    complete = sum(row.complete for row in rows)
    return CombinedCounts(len(rows), complete, len(rows) - complete, len(unreadable))


def _append(sheet: Worksheet, values: Sequence[object]) -> None:
    sheet.append([_text(item) if isinstance(item, str) else item for item in values])


def _percent_text(value: Decimal) -> str:
    return format(value.normalize(), "f")


def _write_result(
    workbook: Workbook,
    parts: Sequence[PartScores],
    subjects: SubjectConfig | None,
    rows: Sequence[_Row],
) -> None:
    sheet = workbook.active
    if sheet is None:
        raise RuntimeError("new workbook must have an active worksheet")
    sheet.title = RESULT_SHEET_NAME
    config = subjects.subjects if subjects is not None else ()
    show_pass = subjects is not None and subjects.criteria.present
    headers = [
        "순번",
        "학번",
        "이름",
        *(f"{part.label} 점수" for part in parts),
        "총점",
        "통합 석차",
        *(subject.name for subject in config),
        *(("합격 여부",) if show_pass else ()),
        "비고",
    ]
    _append(sheet, headers)
    for serial, row in enumerate(rows, 1):
        note = f"확인 필요: {_NOTE_SEPARATOR.join(row.problems)}" if row.problems else ""
        verdict = None if row.passed is None else PASS_TEXT if row.passed else FAIL_TEXT
        _append(
            sheet,
            [
                serial,
                row.student_id,
                row.name,
                *row.part_totals,
                row.total,
                row.rank,
                *row.subject_points,
                *((verdict,) if show_pass else ()),
                note,
            ],
        )
        if not row.complete:
            for cell in sheet[serial + 1]:
                mark_review(cell)
        elif show_pass and row.passed is not None:
            _paint_verdict(sheet.cell(serial + 1, len(headers) - 1), row.passed)
    style_header_row(sheet, _FREEZE_PANES)


def _paint_verdict(cell: StyleableObject, passed: bool) -> None:
    cell.fill = _PASS_FILL if passed else _FAIL_FILL
    cell.font = _PASS_FONT if passed else _FAIL_FONT
    cell.alignment = _CENTER


def _band(sheet: Worksheet, row: int, label: str, width: int) -> None:
    """A gray summary row whose label spans the 순번·학번·이름 columns."""
    sheet.cell(row, 1, label)
    sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
    for column in range(1, width + 1):
        cell = sheet.cell(row, column)
        cell.fill = _BAND_FILL
        cell.font = _BOLD
        cell.alignment = _CENTER


def _average(values: Sequence[Decimal]) -> float | None:
    return float(sum(values, Decimal(0)) / len(values)) if values else None


def _write_subject_results(
    workbook: Workbook,
    parts: Sequence[PartScores],
    subjects: SubjectConfig,
    rows: Sequence[_Row],
) -> None:
    """The first sheet when subjects are configured: one column per subject, the minimum
    scores above the students, below-minimum scores in orange and a red or green verdict."""
    criteria = subjects.criteria
    judged = criteria.present
    per_subject = criteria.per_subject_percent
    total_percent = criteria.total_percent
    config = subjects.subjects
    maxima = [_subject_max(subject, parts) for subject in config]
    total_max = sum((part.maximum for part in parts), Decimal(0))
    sheet = workbook.create_sheet(
        SUBJECT_PASS_SHEET_NAME if judged else SUBJECT_SCORE_SHEET_NAME, 0
    )
    headers = [
        "순번",
        "학번",
        "이름",
        *(subject.name for subject in config),
        "합계",
        *(("미달 과목 수", "판정", "미달 과목") if judged else ()),
        "비고",
    ]
    width = len(headers)
    total_column = 4 + len(config)
    verdict_column = total_column + 2
    _append(sheet, headers)
    for cell in sheet[1]:
        cell.fill = _BAND_FILL
        cell.font = _BOLD
        cell.alignment = _HEADER_ALIGNMENT
    sheet.row_dimensions[1].height = 48

    _band(sheet, 2, "만점", width)
    for column, value in enumerate((*maxima, total_max), 4):
        sheet.cell(2, column, value)
    first = 3
    if judged:
        _band(sheet, 3, "합격 최소 점수", width)
        if per_subject is not None:
            for column, (subject, maximum) in enumerate(zip(config, maxima, strict=True), 4):
                sheet.cell(3, column, _minimum(maximum, per_subject, _whole_points(subject, parts)))
        if total_percent is not None:
            whole = all(_whole_points(subject, parts) for subject in config) and all(
                value == value.to_integral_value()
                for part in parts
                for value in part.question_points
                if value is not None
            )
            sheet.cell(3, total_column, _minimum(total_max, total_percent, whole))
        first = 4

    complete = [row for row in rows if row.complete]
    for serial, row in enumerate(rows, 1):
        line = first + serial - 1
        if not row.complete:
            note = f"확인 필요: {_NOTE_SEPARATOR.join(row.problems)}"
            _append(
                sheet,
                [
                    serial,
                    row.student_id,
                    row.name,
                    *(None for _ in config),
                    None,
                    *((None, REVIEW_TEXT, None) if judged else ()),
                    note,
                ],
            )
            for cell in sheet[line]:
                mark_review(cell)
            continue
        assert row.total is not None
        short = [
            index
            for index, (earned, maximum) in enumerate(zip(row.subject_points, maxima, strict=True))
            if per_subject is not None
            and earned is not None
            and not _reaches(earned, maximum, per_subject)
        ]
        total_short = total_percent is not None and not _reaches(
            row.total, total_max, total_percent
        )
        missed = [config[index].name for index in short]
        if total_short:
            missed.append(TOTAL_SHORT_TEXT)
        verdict = None if row.passed is None else PASS_TEXT if row.passed else FAIL_TEXT
        _append(
            sheet,
            [
                serial,
                row.student_id,
                row.name,
                *row.subject_points,
                row.total,
                *((len(short), verdict, _NOTE_SEPARATOR.join(missed)) if judged else ()),
                "",
            ],
        )
        for index in short:
            sheet.cell(line, 4 + index).fill = _BELOW_FILL
        total_cell = sheet.cell(line, total_column)
        total_cell.font = _BOLD
        if total_short:
            total_cell.fill = _BELOW_FILL
        if judged and row.passed is not None:
            _paint_verdict(sheet.cell(line, verdict_column), row.passed)
            if not row.passed:
                sheet.cell(line, verdict_column + 1).font = Font(
                    name="Calibri", size=11, color="FF9C0006"
                )
        for column in range(1, total_column + 2 if judged else total_column + 1):
            if column != 3:
                sheet.cell(line, column).alignment = _CENTER

    footer = first + len(rows) + 1
    if judged:
        _band(sheet, footer, "과목 통과 인원", width)
        if per_subject is not None:
            for index, maximum in enumerate(maxima):
                sheet.cell(
                    footer,
                    4 + index,
                    sum(
                        _reaches(row.subject_points[index] or Decimal(0), maximum, per_subject)
                        for row in complete
                    ),
                )
        sheet.cell(footer, total_column + 1, "합격 인원")
        sheet.cell(footer, verdict_column, sum(row.passed is True for row in complete))
        footer += 1
    averages = [
        _average([row.subject_points[index] or Decimal(0) for row in complete])
        for index in range(len(config))
    ]
    total_average = _average([row.total or Decimal(0) for row in complete])
    _band(sheet, footer, "과목 평균", width)
    _band(sheet, footer + 1, "과목 평균 정답률", width)
    for column, (average, maximum) in enumerate(
        zip((*averages, total_average), (*maxima, total_max), strict=True), 4
    ):
        mean = sheet.cell(footer, column, average)
        mean.number_format = "0.0"
        rate = sheet.cell(
            footer + 1,
            column,
            None if average is None or not maximum else average / float(maximum),
        )
        rate.number_format = "0%"

    ranges = "; ".join(
        f"{subject.name} "
        + " + ".join(
            f"{parts[item.part - 1].label} {item.first}~{item.last}번" for item in subject.ranges
        )
        for subject in config
    )
    notes = [f"과목 구성: {ranges}"]
    if judged:
        rules = []
        if per_subject is not None:
            rules.append(f"모든 과목 {_percent_text(per_subject)}% 이상")
        if total_percent is not None:
            rules.append(f"총점 {_percent_text(total_percent)}% 이상")
        notes.append(
            f"합격 기준: {' · '.join(rules)}. 주황 = 합격 최소 점수 미달, "
            "판정 초록 = 합격 · 빨강 = 불합격."
        )
    notes.append(
        "노랑 = 확인 필요(학번이 한 파트에만 있거나 확인할 답이 남아 점수를 매기지 않음). "
        "평균과 인원은 점수가 있는 학생만 셉니다."
    )
    line = footer + 3
    for note in notes:
        cell = sheet.cell(line, 1, _text(note))
        cell.font = _NOTE_FONT
        line += 1

    for column, size in ((1, 6), (2, 11), (3, 9)):
        sheet.column_dimensions[get_column_letter(column)].width = size
    for column in range(4, total_column + 1):
        sheet.column_dimensions[get_column_letter(column)].width = 11
    if judged:
        sheet.column_dimensions[get_column_letter(total_column + 1)].width = 8
        sheet.column_dimensions[get_column_letter(verdict_column)].width = 9
        sheet.column_dimensions[get_column_letter(verdict_column + 1)].width = 30
    sheet.column_dimensions[get_column_letter(width)].width = 30
    sheet.freeze_panes = f"D{first}"


def _write_unreadable(
    workbook: Workbook, unreadable: Sequence[tuple[PartScores, PartStudent]]
) -> None:
    sheet = workbook.create_sheet(UNREADABLE_SHEET_NAME)
    _append(sheet, ["파트", "순번", "원본", "점수"])
    for part, student in unreadable:
        _append(sheet, [part.label, student.part_order, student.source_label, student.total])
    style_header_row(sheet, "A2")


def _local_time_text(value: str | None) -> str:
    """A stored UTC ISO time as Korean local ``YYYY-MM-DD HH:MM``; unparseable text stays as is."""
    if not value:
        return ""
    try:
        moment = datetime.fromisoformat(value.replace("Z", "+00:00"))
    except ValueError:
        return value
    if moment.tzinfo is None:
        return value
    return moment.astimezone(_LOCAL_ZONE).strftime("%Y-%m-%d %H:%M")


def _write_parts(workbook: Workbook, parts: Sequence[PartScores], generated_at: str) -> None:
    sheet = workbook.create_sheet(PARTS_SHEET_NAME)
    _append(sheet, ["파트", "시험명", "폴더", "채점 시각", "학생 수", "만점"])
    for part in parts:
        _append(
            sheet,
            [
                part.label,
                part.exam_name,
                part.folder_name,
                _local_time_text(part.graded_at),
                len(part.rows),
                part.maximum,
            ],
        )
    style_header_row(sheet, "A2")
    sheet.append([])
    _append(sheet, ["작성 시각", _local_time_text(generated_at)])


def _write_subjects(
    workbook: Workbook, parts: Sequence[PartScores], subjects: SubjectConfig
) -> None:
    sheet = workbook.create_sheet(SUBJECT_SHEET_NAME)
    _append(sheet, ["과목명", "범위", "만점"])
    for subject in subjects.subjects:
        ranges = ", ".join(
            f"{parts[item.part - 1].label} {item.first}~{item.last}번" for item in subject.ranges
        )
        _append(sheet, [subject.name, ranges, _subject_max(subject, parts)])
    style_header_row(sheet, "A2")
    criteria = subjects.criteria
    if criteria.present:
        sheet.append([])
        _append(sheet, ["합격기준", "기준(%)"])
        for cell in sheet[sheet.max_row]:
            cell.font = Font(name="Calibri", size=11, bold=True)
        if criteria.total_percent is not None:
            _append(sheet, ["총점", criteria.total_percent])
        if criteria.per_subject_percent is not None:
            _append(sheet, ["과목별", criteria.per_subject_percent])


def build_combined_score_book(
    parts: Sequence[PartScores], subjects: SubjectConfig | None, generated_at: str
) -> bytes:
    """Build the 합산 성적표 workbook.

    ``parts`` are in part order (파트1 first); the subject configuration's 파트 numbers count
    from 1 in that order.  Raises ``ValueError`` for inconsistent input; the caller validates
    user-facing problems first.
    """
    _validate(parts, subjects)
    rows, unreadable = _collect(parts, subjects)
    workbook = Workbook()
    _write_result(workbook, parts, subjects, rows)
    if unreadable:
        _write_unreadable(workbook, unreadable)
    _write_parts(workbook, parts, generated_at)
    if subjects is not None:
        _write_subjects(workbook, parts, subjects)
        _write_subject_results(workbook, parts, subjects, rows)
        # Open on the subject sheet, with only that tab selected.
        workbook.active = 0
        for sheet in workbook.worksheets:
            sheet.sheet_view.tabSelected = sheet is workbook.worksheets[0]
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


__all__ = [
    "CombinedCounts",
    "PartScores",
    "PartStudent",
    "build_combined_score_book",
    "summarize_combined_scores",
]
