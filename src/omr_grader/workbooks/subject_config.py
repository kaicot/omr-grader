"""Subject-configuration workbook for combined score reports: strict reader and sample writer.

The workbook has a required ``과목구성`` sheet (과목명, 파트, 시작문항, 끝문항) and an optional
``합격기준`` sheet (항목, 기준(%)).  Like the answer-key reader it is fail-closed: formulas,
oversized packages, and every inconsistent row are refused, and each problem is reported in
Korean in the ``reason`` of its own ``ErrorInfo`` so the teacher can fix the sheet.
"""

from __future__ import annotations

import re
from dataclasses import dataclass
from typing import Any
from decimal import Decimal, InvalidOperation
from io import BytesIO
from pathlib import Path
from zipfile import BadZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from omr_grader.domain.errors import Err, ErrorInfo, Ok, Result
from omr_grader.workbooks.answer_key import _package_bytes

SUBJECT_SHEET_NAME = "과목구성"
CRITERIA_SHEET_NAME = "합격기준"
SUBJECT_HEADERS = ("과목명", "파트", "시작문항", "끝문항")
CRITERIA_HEADERS = ("항목", "기준(%)")
TOTAL_CRITERION = "총점"
PER_SUBJECT_CRITERION = "과목별"
MAX_PARTS = 10
MAX_ROWS = 201
MAX_COLUMNS = 26
MAX_SUBJECTS = 30
MAX_SUBJECT_NAME_LENGTH = 30
MAX_REPORTED_PROBLEMS = 20
# Names the combined sheet already uses for its own columns.
_RESERVED_NAMES = frozenset(
    {
        "순번", "학번", "이름", "총점", "통합 석차", "합격 여부", "비고",
        "합계", "미달 과목 수", "판정", "미달 과목",
    }
)
_PART_TEXT = re.compile(r"(?:파트\s*)?(\d+)")
_PERCENT_TEXT = re.compile(r"(\d+(?:\.\d+)?)\s*%")


@dataclass(frozen=True, slots=True)
class SubjectRange:
    """Questions ``first``..``last`` (inclusive) of part number ``part``."""

    part: int
    first: int
    last: int


@dataclass(frozen=True, slots=True)
class Subject:
    name: str
    ranges: tuple[SubjectRange, ...]


@dataclass(frozen=True, slots=True)
class PassCriteria:
    """Minimum percentages; ``None`` means the criterion is not used."""

    total_percent: Decimal | None = None
    per_subject_percent: Decimal | None = None

    @property
    def present(self) -> bool:
        return self.total_percent is not None or self.per_subject_percent is not None


@dataclass(frozen=True, slots=True)
class SubjectConfig:
    subjects: tuple[Subject, ...]
    criteria: PassCriteria = PassCriteria()

    def part_numbers(self) -> tuple[int, ...]:
        """The 파트 numbers any subject refers to, ascending."""
        return tuple(sorted({item.part for subject in self.subjects for item in subject.ranges}))


def _problem(reason: str, field: str | None = None) -> ErrorInfo:
    return ErrorInfo(
        "SUBJECT_CONFIG_INVALID",
        "error.subject_config_invalid",
        field,
        context={"reason": reason},
    )


def _fail(reason: str, field: str | None = None) -> Err:
    return Err((_problem(reason, field),))


def _whole_number(value: object) -> int | None:
    if isinstance(value, bool):
        return None
    if type(value) is int:
        return value
    if type(value) is float and value.is_integer():
        return int(value)
    return None


def _part_number(value: object) -> int | None:
    number = _whole_number(value)
    if number is None and isinstance(value, str):
        match = _PART_TEXT.fullmatch(value.strip())
        number = int(match.group(1)) if match else None
    return number


def _text(value: object) -> str | None:
    return value.strip() if isinstance(value, str) and value.strip() else None


def _header_columns(
    header: tuple[object, ...], required: tuple[str, ...], sheet: str
) -> dict[str, int] | Err:
    labels = [item.strip() if isinstance(item, str) else None for item in header]
    columns: dict[str, int] = {}
    for name in required:
        if labels.count(name) != 1:
            return _fail(f"'{sheet}' 시트 첫 줄에 '{name}' 열이 하나 있어야 합니다.", sheet)
        columns[name] = labels.index(name)
    return columns


def _data_rows(sheet: Any, width: int) -> list[tuple[int, tuple[object, ...]]]:
    rows = []
    for number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
        cells = tuple(row) + (None,) * max(0, width - len(row))
        if any(item is not None and item != "" for item in cells):
            rows.append((number, cells))
    return rows


def _parse_subjects(sheet: Any) -> Result[tuple[Subject, ...]]:
    max_row = getattr(sheet, "max_row", 0) or 0
    max_column = getattr(sheet, "max_column", 0) or 0
    if max_row > MAX_ROWS or max_column > MAX_COLUMNS:
        return _fail(f"'{SUBJECT_SHEET_NAME}' 시트가 너무 큽니다.", SUBJECT_SHEET_NAME)
    header = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
    columns = _header_columns(tuple(header), SUBJECT_HEADERS, SUBJECT_SHEET_NAME)
    if isinstance(columns, Err):
        return columns
    problems: list[ErrorInfo] = []
    by_name: dict[str, list[SubjectRange]] = {}
    seen: dict[int, list[tuple[int, int, int]]] = {}
    width = max(columns.values()) + 1
    for number, cells in _data_rows(sheet, width):
        where = f"{SUBJECT_SHEET_NAME} {number}행"
        name = _text(cells[columns["과목명"]])
        part = _part_number(cells[columns["파트"]])
        first = _whole_number(cells[columns["시작문항"]])
        last = _whole_number(cells[columns["끝문항"]])
        if name is None:
            problems.append(_problem(f"{where}: 과목명이 비어 있습니다.", where))
            continue
        if len(name) > MAX_SUBJECT_NAME_LENGTH or name in _RESERVED_NAMES:
            problems.append(
                _problem(f"{where}: 과목명 '{name}'은(는) 쓸 수 없는 이름입니다.", where)
            )
            continue
        if part is None or not 1 <= part <= MAX_PARTS:
            problems.append(
                _problem(f"{where}: 파트는 1~{MAX_PARTS} 사이의 숫자(예: 파트1)여야 합니다.", where)
            )
            continue
        if first is None or last is None:
            problems.append(_problem(f"{where}: 시작문항과 끝문항은 숫자여야 합니다.", where))
            continue
        if not (1 <= first <= 100 and 1 <= last <= 100):
            problems.append(_problem(f"{where}: 문항 번호는 1~100 사이여야 합니다.", where))
            continue
        if first > last:
            problems.append(_problem(f"{where}: 시작문항이 끝문항보다 큽니다.", where))
            continue
        clash = next(
            (item for item in seen.get(part, []) if first <= item[1] and item[0] <= last), None
        )
        if clash is not None:
            problems.append(
                _problem(
                    f"{where}: 파트{part}의 {first}~{last}번이 {clash[2]}행의 "
                    f"{clash[0]}~{clash[1]}번과 겹칩니다.",
                    where,
                )
            )
            continue
        seen.setdefault(part, []).append((first, last, number))
        by_name.setdefault(name, []).append(SubjectRange(part, first, last))
    if problems:
        return Err(tuple(problems[:MAX_REPORTED_PROBLEMS]))
    if not by_name:
        return _fail(f"'{SUBJECT_SHEET_NAME}' 시트에 과목이 없습니다.", SUBJECT_SHEET_NAME)
    if len(by_name) > MAX_SUBJECTS:
        return _fail(f"과목은 {MAX_SUBJECTS}개까지 쓸 수 있습니다.", SUBJECT_SHEET_NAME)
    return Ok(
        tuple(
            Subject(name, tuple(sorted(ranges, key=lambda item: (item.part, item.first))))
            for name, ranges in by_name.items()
        )
    )


def _criterion_percent(raw: object, number_format: object) -> Decimal | str:
    """The 기준(%) cell as a percentage, or a Korean reason when it cannot be read safely.

    Typing ``60%`` in Excel stores 0.6 with a percent number format, so a percent-formatted
    number is scaled by 100. An unformatted number up to 1 is ambiguous (0.6 could mean 60% or
    0.6%; a real 1% criterion is implausible) and is refused instead of guessed. ``"60%"`` typed
    as text is accepted.
    """
    if isinstance(raw, str):
        match = _PERCENT_TEXT.fullmatch(raw.strip())
        if match is None:
            return "기준(%)은 숫자여야 합니다."
        return Decimal(match.group(1))
    if type(raw) not in (int, float):
        return "기준(%)은 숫자여야 합니다."
    try:
        percent = Decimal(str(raw))
    except InvalidOperation:
        return "기준(%)은 숫자여야 합니다."
    if not percent.is_finite():
        return "기준(%)은 숫자여야 합니다."
    if isinstance(number_format, str) and "%" in number_format:
        percent = percent * 100
        if percent == percent.to_integral_value():
            percent = percent.quantize(Decimal(1))
    elif 0 < percent <= 1:
        return "기준(%)이 0~1 사이이면 비율인지 퍼센트인지 알 수 없습니다. 60 또는 60%처럼 적으세요."
    return percent


def _parse_criteria(sheet: Any) -> Result[PassCriteria]:
    max_row = getattr(sheet, "max_row", 0) or 0
    max_column = getattr(sheet, "max_column", 0) or 0
    if max_row > MAX_ROWS or max_column > MAX_COLUMNS:
        return _fail(f"'{CRITERIA_SHEET_NAME}' 시트가 너무 큽니다.", CRITERIA_SHEET_NAME)
    header = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
    columns = _header_columns(tuple(header), CRITERIA_HEADERS, CRITERIA_SHEET_NAME)
    if isinstance(columns, Err):
        return columns
    problems: list[ErrorInfo] = []
    found: dict[str, Decimal] = {}
    width = max(columns.values()) + 1
    # Cells (not bare values) so the percent number format of 기준(%) is visible.
    for number, row in enumerate(sheet.iter_rows(min_row=2), 2):
        cells = tuple(row) + (None,) * max(0, width - len(row))
        values = tuple(getattr(cell, "value", None) for cell in cells)
        if not any(item is not None and item != "" for item in values):
            continue
        where = f"{CRITERIA_SHEET_NAME} {number}행"
        item = _text(values[columns["항목"]])
        raw = values[columns["기준(%)"]]
        if item not in (TOTAL_CRITERION, PER_SUBJECT_CRITERION):
            problems.append(_problem(f"{where}: 항목은 '총점' 또는 '과목별'이어야 합니다.", where))
            continue
        if item in found:
            problems.append(_problem(f"{where}: '{item}' 항목이 두 번 나옵니다.", where))
            continue
        percent = _criterion_percent(
            raw, getattr(cells[columns["기준(%)"]], "number_format", None)
        )
        if isinstance(percent, str):
            problems.append(_problem(f"{where}: {percent}", where))
            continue
        if not percent.is_finite() or not 0 <= percent <= 100:
            problems.append(_problem(f"{where}: 기준(%)은 0~100 사이여야 합니다.", where))
            continue
        found[item] = percent
    if problems:
        return Err(tuple(problems[:MAX_REPORTED_PROBLEMS]))
    return Ok(PassCriteria(found.get(TOTAL_CRITERION), found.get(PER_SUBJECT_CRITERION)))


def parse_subject_config_bytes(data: bytes) -> Result[SubjectConfig]:
    """Validate one workbook byte snapshot into a ``SubjectConfig``."""
    package = _package_bytes(data)
    if isinstance(package, Err):
        return _fail(
            "엑셀 파일을 읽을 수 없습니다. 수식, 외부 연결, 매크로가 없는 .xlsx 파일인지 확인하세요."
        )
    workbook = None
    try:
        workbook = load_workbook(
            BytesIO(package.value[0]), read_only=True, data_only=False, keep_links=False
        )
    except (BadZipFile, InvalidFileException, OSError):
        return _fail("엑셀 파일을 열 수 없습니다.")
    try:
        if SUBJECT_SHEET_NAME not in workbook.sheetnames:
            return _fail(f"'{SUBJECT_SHEET_NAME}' 시트가 없습니다.", SUBJECT_SHEET_NAME)
        subjects = _parse_subjects(workbook[SUBJECT_SHEET_NAME])
        if isinstance(subjects, Err):
            return subjects
        criteria: Result[PassCriteria] = Ok(PassCriteria())
        if CRITERIA_SHEET_NAME in workbook.sheetnames:
            criteria = _parse_criteria(workbook[CRITERIA_SHEET_NAME])
        if isinstance(criteria, Err):
            return criteria
        return Ok(SubjectConfig(subjects.value, criteria.value))
    except (OSError, ValueError, KeyError):
        return _fail("엑셀 파일을 읽는 중 문제가 생겼습니다.")
    finally:
        if workbook is not None:
            try:
                workbook.close()
            except OSError:
                pass


def parse_subject_config(path: str) -> Result[SubjectConfig]:
    """Read and validate a subject-config workbook without executing anything in it."""
    try:
        data = Path(path).read_bytes()
    except OSError:
        return _fail("엑셀 파일을 읽을 수 없습니다.", "path")
    return parse_subject_config_bytes(data)


def subject_config_sample_bytes() -> bytes:
    """A deterministic sample: a two-part graduation exam of seven subjects with pass criteria."""
    workbook = Workbook()
    sheet = workbook.active
    if sheet is None:
        raise RuntimeError("new workbook must have an active worksheet")
    sheet.title = SUBJECT_SHEET_NAME
    sheet.append(list(SUBJECT_HEADERS))
    for row in (
        ("해부생리학", 1, 1, 30),
        ("근골격계 작업치료", 1, 31, 60),
        ("신경계 작업치료", 1, 61, 90),
        ("보건의료관계법규", 1, 91, 100),
        ("보건의료관계법규", 2, 1, 10),
        ("작업치료평가", 2, 11, 40),
        ("아동작업치료", 2, 41, 70),
        ("정신사회작업치료", 2, 71, 100),
    ):
        sheet.append(list(row))
    criteria = workbook.create_sheet(CRITERIA_SHEET_NAME)
    criteria.append(list(CRITERIA_HEADERS))
    criteria.append([TOTAL_CRITERION, 60])
    criteria.append([PER_SUBJECT_CRITERION, 60])
    guide = workbook.create_sheet("설명")
    for line in (
        "과목구성: 과목마다 파트와 문항 범위를 적습니다. 같은 과목명을 여러 줄에 적으면 범위가 합쳐집니다.",
        "파트는 합산할 시험을 고른 순서대로 1, 2, ...입니다. 같은 파트 안에서 범위가 겹치면 안 됩니다.",
        "합격기준(선택): 총점 = 전체 만점 대비 최소 %, 과목별 = 모든 과목에서 필요한 최소 %.",
        "기준(%)은 60 또는 60%처럼 적습니다. 0.6처럼 비율로 적으면 퍼센트인지 알 수 없어 읽지 않습니다.",
        "합격기준 시트나 그 안의 줄은 없어도 됩니다. 없는 기준은 적용하지 않습니다.",
        "이 시트와 표 오른쪽의 다른 열은 읽지 않습니다.",
    ):
        guide.append([line])
    guide.column_dimensions["A"].width = 100
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


__all__ = [
    "CRITERIA_SHEET_NAME",
    "PER_SUBJECT_CRITERION",
    "SUBJECT_SHEET_NAME",
    "TOTAL_CRITERION",
    "PassCriteria",
    "Subject",
    "SubjectConfig",
    "SubjectRange",
    "parse_subject_config",
    "parse_subject_config_bytes",
    "subject_config_sample_bytes",
]
